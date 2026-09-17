"""Logic campaign coverage (A09/A11): RS29511/RS1G08/RS0204 sheet_map vs specs.

Run: python -m ate.core.check_logic_campaign
"""
from __future__ import annotations

import sys
from pathlib import Path

import re
import yaml
from openpyxl import load_workbook

from ate.core.database import find_campaign_root, set_context
from ate.core.paths import PARTS_DIR, TEST_DB_ROOT
from ate.core.registry import all_tests, load_family
from ate.fixture.modes import enabled_tests_for_part, logic_catalog_for_ui

_NON_TEST = frozenset({"Summary", "Checklist"})


def _workbook_sheets(path: Path) -> set[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return {s for s in wb.sheetnames if s not in _NON_TEST}
    finally:
        wb.close()


def _check_part(
    part_key: str,
    component: str,
    part: str,
    package: str,
    version: str,
    *,
    operator: str,
) -> list[str]:
    errors: list[str] = []
    yaml_path = PARTS_DIR / f"{part_key}.yaml"
    if not yaml_path.is_file():
        return [f"part yaml missing: {yaml_path}"]

    ctx = set_context(
        component=component,
        part=part,
        package=package,
        operator=operator,
        version=version,
        part_key=part_key,
    )
    load_family("logic")
    enabled = enabled_tests_for_part(part_key, catalog=ctx.load_test_catalog())
    if not enabled:
        errors.append(f"{part_key}: enabled_tests empty")
        return errors

    ids = {t.id for t in all_tests()}
    missing_ids = [t for t in enabled if t not in ids]
    if missing_ids:
        errors.append(f"{part_key}: enabled ids not registered: {missing_ids}")

    modes = logic_catalog_for_ui(part_key)
    if not modes or modes[0].get("mode") != "LOGIC":
        errors.append(f"{part_key}: logic_catalog_for_ui must expose LOGIC")
    if any(m.get("mode") in ("G11", "BUFFER", "G_NEG100") for m in modes):
        errors.append(f"{part_key}: OPA fixture modes leaked into Logic catalog")

    sm_path = ctx.sheet_map_path()
    if not sm_path.is_file():
        errors.append(f"{part_key}: sheet_map missing {sm_path}")
        return errors
    data = yaml.safe_load(sm_path.read_text(encoding="utf-8")) or {}
    tests = data.get("tests") or {}
    allow = set(enabled)
    sheets = {t.lab_sheet: t.id for t in all_tests() if t.lab_sheet and t.id in allow}
    for key, entry in tests.items() if isinstance(tests, dict) else []:
        if not isinstance(entry, dict):
            continue
        sheet = entry.get("excel_sheet")
        if sheet and sheet not in sheets:
            errors.append(f"{part_key}: sheet_map {key}->{sheet} has no enabled lab_sheet")

    xlsx = ctx.lab_report_path()
    if not xlsx.is_file():
        errors.append(f"{part_key}: workbook missing {xlsx}")
    else:
        wb = _workbook_sheets(xlsx)
        map_sheets = {
            str(e.get("excel_sheet"))
            for e in (tests.values() if isinstance(tests, dict) else [])
            if isinstance(e, dict) and e.get("excel_sheet")
        }
        missing_wb = sorted(map_sheets - wb)
        if missing_wb:
            errors.append(f"{part_key}: workbook missing sheets {missing_wb}")
        if part_key == "rs0204":
            unmapped = sorted(wb - map_sheets)
            if unmapped:
                errors.append(f"{part_key}: workbook sheets not in sheet_map: {unmapped}")

    # Restore OpAmp default so other checks stay stable
    set_context(
        component="OpAmp",
        part="RS622",
        package="TTSOP8",
        operator="Eugene",
        version="Version_1",
        part_key="rs622",
    )
    load_family("opamp")
    return errors


def check_logic_campaign() -> list[str]:
    errors: list[str] = []
    need = [
        ("Logic", "RS29511", "SOIC", "Version_1", "Soo"),
        ("Logic", "RS1G08", "SOT23", "Version_1", "Ariff"),
        ("Logic", "RS0204", "TSSOP14", "Version_1", "ChangThong"),
    ]
    for component, part, package, version, operator in need:
        root = find_campaign_root(component, part, package, version, operator=operator)
        if root is None:
            errors.append(f"campaign missing: {component}/{part}/{package}/[{operator}|legacy]/{version}")
    if errors:
        return errors

    errors += _check_part("rs29511", "Logic", "RS29511", "SOIC", "Version_1", operator="Soo")
    errors += _check_part("rs1g08", "Logic", "RS1G08", "SOT23", "Version_1", operator="Ariff")
    errors += _check_part("rs0204", "Logic", "RS0204", "TSSOP14", "Version_1", operator="ChangThong")

    # Probe Ariff + RS0204 ids registered when Logic loads
    load_family("logic")
    ids = {t.id for t in all_tests()}
    for need in (
        "tp",
        "delta_supply_current",
        "off_current",
        "input_thresholds",
        "ioff_leakage",
        "input_leakage_sweep",
        "supply_current_sweep",
        "vih_vil",
        "voh_load",
        "vol_load",
        "input_threshold",
        "delta_icc",
        "ii",
        "ioz",
        "vih",
        "tp_rs0204",
        "tpd",
    ):
        if need not in ids:
            errors.append(f"logic registry missing {need}")

    # Soo catalog must not leak Ariff-only A12 ids
    soo_en = set(enabled_tests_for_part("rs29511") or [])
    for banned in (
        "delta_supply_current",
        "supply_current_sweep",
        "vih_vil",
        "voh_load",
        "vol_load",
        "ioff_leakage",
        "input_leakage_sweep",
        "off_current",
    ):
        if banned in soo_en:
            errors.append(f"rs29511 must not enable Ariff-only id {banned}")

    ariff_en = set(enabled_tests_for_part("rs1g08") or [])
    for need in ("vih_vil", "voh_load", "vol_load", "supply_current_sweep"):
        if need not in ariff_en:
            errors.append(f"rs1g08 enabled_tests missing {need}")

    ariff_src = Path(__file__).resolve().parents[1] / "tests" / "logic" / "ariff_dc.py"
    ariff_text = ariff_src.read_text(encoding="utf-8")
    if re.search(r"(?<![\"'\w])input\s*\(", ariff_text):
        errors.append("ariff_dc.py must not call input()")
    # Real import only (ignore docs that say "do not import Ariff")
    if re.search(r"(?m)^\s*(?:import\s+Ariff\b|from\s+Ariff\b)", ariff_text):
        errors.append("ariff_dc.py must not import Ariff.*")

    from ate.tests.logic.wraps import _invoke_legacy
    from ate.core.runner import RunParams as _RP

    class _I:
        pass

    def _one(_instr, vcc):
        return {"VCC": vcc}

    def _two(_instr, vcca, vccb):
        return {"VCCA": vcca, "VCCB": vccb}

    p = _RP(vcc=1.8, part="rs29511", current_limit_a=0.05)
    one = _invoke_legacy(_one, _I(), p)
    if one.get("VCC") != 1.8:
        errors.append(f"wraps one-rail IDD got {one}")
    two = _invoke_legacy(_two, _I(), p)
    if two.get("VCCA") != 1.8 or two.get("VCCB") != 1.8:
        errors.append(f"wraps dual-rail IDD without yaml vccb got {two}")
    p_rs = _RP(vcc=1.8, part="rs0204", current_limit_a=0.05)
    two_rs = _invoke_legacy(_two, _I(), p_rs)
    if two_rs.get("VCCB") != 3.3:
        errors.append(f"wraps dual-rail IDD rs0204 yaml vccb got {two_rs}")
    from ate.core.registry import get as _get
    from ate.core.runner import RunParams
    from ate.tests.logic.rs0204 import rails_from_params
    from pathlib import Path as _P

    rs_src = _P(__file__).resolve().parents[1] / "tests" / "logic" / "rs0204.py"
    if "Soo.logic_tests" in rs_src.read_text(encoding="utf-8"):
        errors.append("rs0204.py must not import Soo.logic_tests")
    try:
        rails_from_params(RunParams(vcc=1.8, current_limit_a=0.05))
    except Exception as exc:
        errors.append(f"rs0204 rails_from_params(1.8) failed: {exc}")
    try:
        rails_from_params(RunParams(vcc=5.0, current_limit_a=0.05))
        errors.append("rs0204 rails_from_params must refuse VCCA 5 > VCCB")
    except RuntimeError:
        pass
    spec = _get("vih")
    if spec is not None:
        try:
            spec.run(None, None)
            errors.append("rs0204 vih must require instruments")
        except RuntimeError as exc:
            if "Missing instruments" not in str(exc):
                errors.append(f"rs0204 vih should say missing instruments, got {exc!r}")
        except Exception as exc:
            errors.append(f"rs0204 vih unexpected {type(exc).__name__}: {exc}")
    errors += _path_b_lock_campaigns()
    load_family("opamp")
    return errors


def _path_b_lock_campaigns() -> list[str]:
    """Path B lock: second golden xlsx under Version workbook/ is FAIL.

    ultimate_manual jot books are never_auto_write and are not Version orphans.
    """
    from ate.tests.logic.excel_lock import golden_xlsx, list_xlsx, uses_excel_lock
    from ate.tests.logic.product_model import has_product_model, load_product_model

    errors: list[str] = []
    if not TEST_DB_ROOT.is_dir():
        return errors
    for wb_dir in TEST_DB_ROOT.glob("Logic/*/*/*/Version_*/workbook"):
        try:
            rel = wb_dir.relative_to(TEST_DB_ROOT)
        except ValueError:
            continue
        parts = rel.parts
        if len(parts) < 2:
            continue
        key = str(parts[1]).strip().lower()
        if not has_product_model(key):
            continue
        model = load_product_model(key)
        if not uses_excel_lock(model):
            continue
        gold = golden_xlsx(wb_dir)
        names = [p.name for p in gold]
        if len(gold) > 1:
            errors.append(
                f"{key}: workbook/ has orphan golden xlsx {sorted(names)} "
                "(golden_auto one_per_version_overwrite; ultimate_manual excluded)"
            )
        all_names = [p.name for p in list_xlsx(wb_dir)]
        if any(n.endswith("_filled.xlsx") or Path(n).stem.endswith("_filled") for n in all_names):
            errors.append(f"{key}: Path B must not keep _filled.xlsx under workbook/")
    return errors


def main() -> int:
    errors = check_logic_campaign()
    if errors:
        print("FAIL logic-campaign:")
        for line in errors:
            print(f"  - {line}")
        return 1
    print(
        "OK logic-campaign: RS29511 + RS1G08 + RS0204 sheet_map / enabled_tests / "
        "LOGIC fixture / A12 Ariff ids"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
