"""Copy an existing lab xlsx into a campaign workbook/ and stub sheet_map.

Paste cells are never guessed — stubs use FILL_ME for the operator.
A sheet_map that already has real Excel anchors is left intact (backed
up first if we must rewrite a stub).
"""
from __future__ import annotations

import re
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Optional

import yaml

MYT = timezone(timedelta(hours=8))
_CELL_RE = re.compile(r"^[A-Z]{1,3}\d{1,5}$")
_FILL_TOKENS = {"FILL_ME", "STUB", "TODO", "TBD", ""}


def _folder_key(sheet: str) -> str:
    raw = re.sub(r"[^\w\-]+", "", (sheet or "").replace(" ", ""))
    return raw or "Sheet"


def _is_real_cell(value: Any) -> bool:
    if isinstance(value, str):
        token = value.split("#", 1)[0].strip()
        if token.upper() in _FILL_TOKENS:
            return False
        return bool(_CELL_RE.match(token))
    if isinstance(value, (list, tuple)):
        return any(_is_real_cell(v) for v in value)
    if isinstance(value, dict):
        return any(_is_real_cell(v) for v in value.values())
    return False


def sheet_map_has_paste_anchors(data: dict[str, Any]) -> bool:
    tests = data.get("tests")
    if not isinstance(tests, dict) or not tests:
        return False
    for entry in tests.values():
        if not isinstance(entry, dict):
            continue
        paste = entry.get("paste")
        if _is_real_cell(paste):
            return True
    return False


def _backup(path: Path) -> Optional[Path]:
    if not path.is_file():
        return None
    ts = datetime.now(MYT).strftime("%Y%m%d_%H%M%S")
    dest = path.with_name(f"{path.stem}.bak_{ts}{path.suffix}")
    shutil.copy2(path, dest)
    return dest


def _sheet_names(xlsx: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _pick_xlsx_dialog() -> str:
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    try:
        root.attributes("-topmost", True)
    except Exception:
        pass
    path = filedialog.askopenfilename(
        title="Select existing lab workbook (.xlsx)",
        filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")],
    )
    root.destroy()
    return str(path or "")


def _stub_sheet_map(ctx, dest_name: str, sheets: list[str]) -> dict[str, Any]:
    tests: dict[str, Any] = {}
    used: set[str] = set()
    for name in sheets:
        key = _folder_key(name)
        base = key
        n = 2
        while key in used:
            key = f"{base}{n}"
            n += 1
        used.add(key)
        tests[key] = {
            "folder": key,
            "excel_sheet": name,
            "fixture_mode": "FILL_ME",
            "automated": False,
            "dut_iterations": int(ctx.sample_size or 4),
            "paste": {
                "results": "FILL_ME",
                "setup_text": "FILL_ME",
                "photos": "FILL_ME",
            },
        }
    return {
        "component": ctx.component,
        "part": ctx.part,
        "package": ctx.package,
        "version": ctx.version,
        "sample_size": int(ctx.sample_size or 4),
        "workbook": {"path": f"../workbook/{dest_name}"},
        "naming": {
            "screenshot": "{TEST}_{DUT}_{VARIANT}_{TIMESTAMP}.jpg",
            "graph": "{TEST}_{DUT}_{VARIANT}_{TIMESTAMP}.png",
        },
        "tests": tests,
    }


def _write_stub_yaml(path: Path, data: dict[str, Any]) -> None:
    header = (
        "# STUB sheet_map — scaffolded from workbook sheet names.\n"
        "# Replace every FILL_ME with a real Excel anchor (e.g. B3).\n"
        "# Do not invent paste cells; leave FILL_ME until the operator measures the sheet.\n"
    )
    body = yaml.safe_dump(data, sort_keys=False, allow_unicode=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(header + body, encoding="utf-8")


def import_workbook(
    *,
    source_path: str = "",
    dest_name: str = "",
    pick: bool = False,
) -> dict[str, Any]:
    """Copy xlsx into the active campaign workbook/ and stub sheet_map if needed."""
    from ate.core.database import get_context

    src_str = str(source_path or "").strip()
    if pick or not src_str:
        src_str = _pick_xlsx_dialog()
    if not src_str:
        raise ValueError("No workbook selected")

    src = Path(src_str).expanduser()
    if not src.is_file():
        raise FileNotFoundError(f"Workbook not found: {src}")
    if src.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx workbooks can be imported")

    name = Path(dest_name).name if dest_name else src.name
    if not name.lower().endswith(".xlsx"):
        name = f"{name}.xlsx"
    if not name or name.startswith("."):
        raise ValueError("Invalid destination workbook name")

    ctx = get_context()
    ctx.ensure_tree()
    wb_dir = ctx.workbook_dir()
    wb_dir.mkdir(parents=True, exist_ok=True)
    dest = wb_dir / name

    xlsx_backup = None
    copied = True
    if dest.resolve() == src.resolve():
        copied = False
    else:
        if dest.is_file():
            xlsx_backup = _backup(dest)
        shutil.copy2(src, dest)

    sheets = _sheet_names(dest)
    sm_path = ctx.sheet_map_path()
    existing = ctx.load_sheet_map()
    map_backup = None
    if existing and sheet_map_has_paste_anchors(existing):
        action = "left_intact"
        note = "Existing sheet_map has paste anchors — not overwritten."
    elif existing:
        map_backup = _backup(sm_path)
        _write_stub_yaml(sm_path, _stub_sheet_map(ctx, name, sheets))
        action = "stub_replaced"
        note = "Previous stub/empty sheet_map backed up; new stub written from sheet names."
    else:
        _write_stub_yaml(sm_path, _stub_sheet_map(ctx, name, sheets))
        action = "stub_created"
        note = "Stub sheet_map written from sheet names (FILL_ME paste cells)."

    created = ctx.ensure_tree()
    return {
        "ok": True,
        "source": str(src),
        "workbook": str(dest),
        "copied": copied,
        "xlsx_backup": str(xlsx_backup) if xlsx_backup else None,
        "sheet_map": str(sm_path),
        "sheet_map_action": action,
        "sheet_map_backup": str(map_backup) if map_backup else None,
        "sheets": sheets,
        "note": note,
        "context": ctx.identity(),
        "created": created,
    }
