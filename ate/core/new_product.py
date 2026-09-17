"""New product under test: RUN-IC class -> campaign folders + DEMO mock walk.

Does not scrape en.run-ic.com. Catalog SKUs are added only when we test them.
"""
from __future__ import annotations

import json
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Optional

import yaml

from ate.core.paths import CONFIG_DIR, PARTS_DIR, TEST_DB_ROOT

MYT = timezone(timedelta(hours=8))
_PART_KEY_RE = re.compile(r"[^a-z0-9]+")


def load_categories() -> list[dict[str, Any]]:
    path = CONFIG_DIR / "run_ic.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    rows = data.get("categories") if isinstance(data, dict) else None
    return [r for r in rows if isinstance(r, dict) and r.get("id")] if isinstance(rows, list) else []


def load_inventory() -> list[dict[str, Any]]:
    path = CONFIG_DIR / "inventory.yaml"
    if not path.is_file():
        return []
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    rows = data.get("parts") if isinstance(data, dict) else None
    return [r for r in rows if isinstance(r, dict) and r.get("part")] if isinstance(rows, list) else []


SHEET_CLASS_CATEGORY = {
    "Low Noise Op-Amp": "opamp",
    "General Op-Amp": "opamp",
    "Precision Op-Amp": "opamp",
    "Analog Switch": "analog_switch",
    "Level Shifters": "level",
    "Logic Series": "logic",
    "Linear Regulator": "power",
}


def _inventory_meta() -> dict[str, Any]:
    path = CONFIG_DIR / "inventory.yaml"
    if not path.is_file():
        return {}
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return data if isinstance(data, dict) else {}


def reports_root() -> Path:
    raw = str(_inventory_meta().get("reports_root") or "").strip()
    return Path(raw) if raw else Path()


def qualification_xlsx() -> Path | None:
    meta = _inventory_meta()
    for key in ("qualification_xlsx",):
        raw = str(meta.get(key) or "").strip()
        if raw:
            p = Path(raw)
            if p.is_file():
                return p
    root = reports_root()
    if root.is_dir():
        p = root / "Qualification Product List-20260320.xlsx"
        if p.is_file():
            return p
    return None


def load_qualification_main() -> list[dict[str, Any]]:
    """Main sheet SKUs from the tracking-list xlsx. Empty if the file is not on this PC."""
    path = qualification_xlsx()
    if path is None:
        return []
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Main" not in wb.sheetnames:
            return []
        out: list[dict[str, Any]] = []
        for row in wb["Main"].iter_rows(min_row=3, values_only=True):
            sheet_class = str(row[1] or "").strip() if row else ""
            part = str(row[2] or "").strip() if row else ""
            if not part or not sheet_class:
                continue
            out.append(
                {
                    "sheet_class": sheet_class,
                    "part": part,
                    "model": str(row[3] or "").strip() if row else "",
                    "package": str(row[4] or "").strip() if row else "",
                    "lot": str(row[5] or "").strip() if row else "",
                    "pic": str(row[7] or "").strip() if row else "",
                    "status": str(row[9] or "").strip() if row else "",
                }
            )
        return out
    finally:
        wb.close()


def resolve_inventory_report(inv: dict[str, Any] | None) -> Path | None:
    """First real xlsx for this tracking row (file or folder under reports_root)."""
    rel = str((inv or {}).get("report") or "").strip()
    if not rel:
        return None
    src = Path(rel)
    if not src.is_absolute():
        root = reports_root()
        src = (root / rel) if str(root) else src
    if src.is_file() and src.suffix.lower() == ".xlsx":
        return src
    if src.is_dir():
        xlsxs = [
            p
            for p in src.iterdir()
            if p.is_file()
            and p.suffix.lower() == ".xlsx"
            and not p.name.upper().startswith("[OUTDATED]")
        ]
        if not xlsxs:
            return None
        part = str((inv or {}).get("part") or "").upper()
        named = [p for p in xlsxs if part and part in p.name.upper()]
        pick = named[0] if named else xlsxs[0]
        return pick
    return None


def _maybe_merge_golden_style(ingested: dict[str, Any]) -> dict[str, Any]:
    """Apply ORT/golden culture only when the copied workbook has those sheets.

    Skip when GBW/VOS already hold numbers or formulas -- golden merge-center
    would cover the DUT result cells (SOP8 tracking ingest).
    """
    path = Path(str(ingested.get("workbook") or ""))
    if not path.is_file():
        return ingested
    try:
        from openpyxl import load_workbook

        from ate.reporting.golden_layout import TEST_SHEETS
        from ate.reporting.golden_workbook import apply_golden_workbook

        wb = load_workbook(path, read_only=True)
        names = set(wb.sheetnames)
        filled = False
        probes = {
            "GBW": ("C21", "R20", "F15"),
            "VOS": ("B16", "R16"),
        }
        for sheet, coords in probes.items():
            if sheet not in names:
                continue
            ws = wb[sheet]
            for coord in coords:
                v = ws[coord].value
                if v not in (None, ""):
                    filled = True
                    break
            if filled:
                break
        wb.close()
        if not (names & set(TEST_SHEETS)):
            ingested["golden"] = False
            ingested["golden_note"] = "No OpAmp golden sheets in this workbook -- layout left as copied."
            return ingested
        if filled:
            ingested["golden"] = False
            ingested["golden_note"] = "Tracking workbook already has GBW/VOS numbers; golden merge skipped."
            return ingested
        apply_golden_workbook(path)
        ingested["golden"] = True
    except Exception as exc:
        ingested["golden"] = False
        ingested["golden_error"] = str(exc)
    return ingested


def category_by_id(cat_id: str) -> dict[str, Any]:
    key = (cat_id or "").strip().lower()
    for row in load_categories():
        if str(row.get("id") or "").lower() == key:
            return row
        if str(row.get("run_ic") or "").lower() == key:
            return row
        for alias in row.get("sheet_aliases") or []:
            if str(alias).lower() == key:
                return row
    raise ValueError(f"Unknown RUN-IC category {cat_id!r}")


def part_key_for(part: str) -> str:
    return _PART_KEY_RE.sub("", (part or "").strip().lower())


def _inventory_match(part: str, package: str = "", model: str = "") -> dict[str, Any] | None:
    """Best inventory row for this SKU (package/model break ties)."""
    part_u = (part or "").strip().upper()
    if not part_u:
        return None
    rows = [r for r in load_inventory() if str(r.get("part") or "").upper() == part_u]
    if not rows:
        return None
    pkg = (package or "").strip()
    model_u = (model or "").strip().upper()

    def _score(r: dict[str, Any]) -> int:
        s = 0
        if pkg and str(r.get("package") or "") == pkg:
            s += 2
        if model_u and str(r.get("model") or "").upper() == model_u:
            s += 2
        return s

    rows.sort(key=_score, reverse=True)
    return rows[0]


def suite_for_part(
    part: str,
    *,
    component: str = "",
    package: str = "",
    model: str = "",
) -> str:
    """Test family to load. Folder class can be Level while ate_suite is logic."""
    from ate.core.database import family_for_component

    inv = _inventory_match(part, package, model)
    ate = str((inv or {}).get("ate_suite") or "").strip()
    if ate:
        return ate
    return family_for_component(component)


def _test_track_rows(pk: str, family: str) -> list[tuple[str, str]]:
    """enabled test id -> folder/excel name. Does not invent paste cells."""
    from ate.fixture.modes import enabled_tests_for_part
    from ate.core.registry import get as reg_get
    from ate.core.registry import load_family

    ids = list(enabled_tests_for_part(pk) or [])
    if not ids:
        return []
    if family:
        try:
            load_family(family)
        except Exception:
            pass
    rows: list[tuple[str, str]] = []
    for tid in ids:
        spec = reg_get(tid)
        folder = str(getattr(spec, "lab_sheet", "") or tid)
        rows.append((tid, folder))
    return rows


def _tracked_sheet_map(
    cat: dict[str, Any],
    part: str,
    package: str,
    operator: str,
    pk: str,
    family: str,
) -> str:
    rows = _test_track_rows(pk, family)
    lines = [
        f"component: {cat['component']}",
        f"part: {part}",
        f"package: {package}",
        f"operator: {operator}",
        f"version: Version_1",
        "tests:",
    ]
    if not rows:
        lines.append("  {}")
    else:
        for tid, folder in rows:
            lines.append(f"  {tid}:")
            lines.append(f"    folder: {folder}")
            lines.append(f"    excel_sheet: {folder}")
    return "\n".join(lines) + "\n"


def _tracked_catalog(pk: str) -> str:
    from ate.fixture.modes import enabled_tests_for_part

    ids = list(enabled_tests_for_part(pk) or [])
    if not ids:
        return "enabled_tests: []\n"
    return "enabled_tests:\n" + "".join(f"  - {i}\n" for i in ids)


def _ensure_test_folders(root: Path, folders: list[str], sample: int) -> list[str]:
    created: list[str] = []
    n = max(1, int(sample))
    for key in folders:
        if not key:
            continue
        for dut in range(1, n + 1):
            for sub in ("screenshots", "graphs"):
                p = root / key / f"DUT_{dut}" / sub
                if not p.exists():
                    p.mkdir(parents=True, exist_ok=True)
                    created.append(str(p))
        shared = root / key / "screenshots"
        if not shared.exists():
            shared.mkdir(parents=True, exist_ok=True)
            created.append(str(shared))
    return created


def campaign_root(
    component: str,
    part: str,
    package: str,
    version: str = "Version_1",
    *,
    operator: str = "Eugene",
    base: Optional[Path] = None,
) -> Path:
    from ate.core.database import require_write_operator

    op = require_write_operator(operator)
    return (base or TEST_DB_ROOT) / component / part / package / op / version


def _write_if_missing(path: Path, text: str) -> bool:
    if path.is_file():
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return True


def _stub_part_yaml(cat: dict[str, Any], part: str, package: str, model: str, sample: int) -> str:
    pk = part_key_for(part)
    fam = str(cat.get("family") or "")
    mode = "LDO" if fam == "power" else ("LOGIC" if fam in ("logic", "switch", "level") else "BUFFER")
    return (
        f"# Stub -- fill enabled_tests when recipes exist. Do not import vendor trees.\n"
        f"part: {part}\n"
        f"model: {model or part}\n"
        f"component: {cat['component']}\n"
        f"package: {package}\n"
        f"version: Version_1\n"
        f"sample_size: {int(sample)}\n"
        f"year: \"{datetime.now(MYT).year}\"\n"
        f"part_key: {pk}\n"
        f"product_class: {cat['id']}\n"
        f"vcc: 5.0\n"
        f"enabled_tests: []\n"
        f"fixture_modes:\n"
        f"  {mode}:\n"
        f"    label: \"{cat.get('run_ic') or cat['id']} board ({part})\"\n"
        f"    board_class: general\n"
        f"    tests: []\n"
    )


def _stub_sheet_map(cat: dict[str, Any], part: str, package: str, operator: str) -> str:
    return (
        f"component: {cat['component']}\n"
        f"part: {part}\n"
        f"package: {package}\n"
        f"operator: {operator}\n"
        f"version: Version_1\n"
        f"tests: {{}}\n"
    )


def next_version_name(existing: list[str] | None = None) -> str:
    """Return next Version_N from existing campaign folder names."""
    from ate.core.database import _VERSION_RE

    nums: list[int] = []
    for name in existing or []:
        m = _VERSION_RE.match(str(name or ""))
        if m:
            nums.append(int(m.group(1)))
    return f"Version_{max(nums) + 1 if nums else 1}"


def list_operator_versions(
    component: str,
    part: str,
    package: str,
    operator: str,
    *,
    base: Optional[Path] = None,
) -> list[str]:
    from ate.core.database import is_campaign_dir, require_write_operator

    op = require_write_operator(operator)
    folder = (base or TEST_DB_ROOT) / component / part / package / op
    if not folder.is_dir():
        return []
    versions = [p.name for p in sorted(folder.iterdir()) if p.is_dir() and is_campaign_dir(p)]
    return versions


def ensure_version(
    *,
    component: str = "",
    part: str = "",
    package: str = "",
    operator: str = "",
    version: str = "",
    sample_size: int = 4,
    copy_from_version: str = "",
    open_folder: bool = False,
    apply: bool = True,
    base: Optional[Path] = None,
) -> dict[str, Any]:
    """Create next (or named) Version_N under operator. Does not clone xlsx."""
    from ate.core.database import get_context, remember_operator, set_context

    ctx = get_context()
    component = str(component or ctx.component or "").strip() or "OpAmp"
    part = str(part or ctx.part or "").strip().upper()
    package = str(package or ctx.package or "").strip() or "SOT23"
    if not part:
        raise ValueError("part is required")
    op = remember_operator(
        operator or ctx.operator,
        component=component,
        part=part,
        package=package,
    )
    existing = list_operator_versions(component, part, package, op, base=base)
    ver = str(version or "").strip() or next_version_name(existing)
    if ver in existing:
        raise ValueError(f"{ver} already exists under {op}")
    sample = max(1, int(sample_size or ctx.sample_size or 4))
    root = campaign_root(component, part, package, ver, operator=op, base=base)
    created: list[str] = []
    for rel in ("_manifest", "sessions", "workbook"):
        p = root / rel
        if not p.exists():
            p.mkdir(parents=True, exist_ok=True)
            created.append(str(p))
    for dut in range(1, sample + 1):
        for sub in ("screenshots", "graphs"):
            p = root / "Setup" / f"DUT_{dut}" / sub
            if not p.exists():
                p.mkdir(parents=True, exist_ok=True)
                created.append(str(p))

    src_ver = str(copy_from_version or ctx.version or "").strip()
    src_root = None
    if src_ver and src_ver != ver:
        cand = campaign_root(component, part, package, src_ver, operator=op, base=base)
        if cand.is_dir():
            src_root = cand
    for name in ("sheet_map.yaml", "test_catalog.yaml"):
        dest = root / "_manifest" / name
        if dest.is_file():
            continue
        if src_root is not None:
            src = src_root / "_manifest" / name
            if src.is_file():
                dest.write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
                created.append(str(dest))
                continue
        if name == "sheet_map.yaml":
            dest.write_text(
                f"component: {component}\npart: {part}\npackage: {package}\n"
                f"operator: {op}\nversion: {ver}\ntests: {{}}\n",
                encoding="utf-8",
            )
        else:
            dest.write_text("enabled_tests: []\n", encoding="utf-8")
        created.append(str(dest))

    applied = None
    if apply and base is None:
        applied = set_context(
            component=component,
            part=part,
            package=package,
            operator=op,
            version=ver,
            model=str(ctx.model or part),
            part_key=part_key_for(part),
            sample_size=sample,
        ).identity()
    if open_folder and base is None:
        try:
            os.startfile(str(root))
        except OSError:
            pass
    return {
        "ok": True,
        "root": str(root),
        "version": ver,
        "operator": op,
        "created": created,
        "versions": list_operator_versions(component, part, package, op, base=base),
        "context": applied,
        "note": "Version folder created. Import xlsx separately if needed.",
    }


def new_run_session(*, run_label: str = "", params: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    """Write a new sessions/session_*.json run record. Does not open VISA."""
    from ate.core.database import begin_session, get_context

    ctx = get_context()
    ctx.ensure_tree()
    p = dict(params or {})
    label = str(run_label or p.get("run_label") or "").strip()
    if label:
        p["run_label"] = label
    p.setdefault("manual_session", True)
    session = begin_session(p, instrument_map={}, fixture_plan=[])
    path = ctx.sessions_dir() / f"{session.session_id}.json"
    return {
        "ok": True,
        "session_id": session.session_id,
        "path": str(path),
        "context": ctx.identity(),
        "note": "Run record created. Discover / Open Session still required for instruments.",
    }


def ensure_product(
    *,
    category_id: str,
    part: str,
    package: str,
    model: str = "",
    sample_size: int = 4,
    operator: str = "",
    open_folder: bool = True,
    apply: bool = True,
    base: Optional[Path] = None,
) -> dict[str, Any]:
    """Create campaign tree + stub yaml. Does not overwrite filled maps."""
    from ate.core.database import remember_operator

    part = str(part or "").strip().upper()
    package = str(package or "").strip() or "SOT23"
    if not part:
        raise ValueError("part is required")
    cat = category_by_id(category_id)
    inv = _inventory_match(part, package, model)
    ate_suite = str((inv or {}).get("ate_suite") or "").strip()
    # Folder follows tracking class (Level / Power). ate_suite only loads tests.
    component = str(cat["component"])
    op = remember_operator(
        operator,
        component=component,
        part=part,
        package=package,
        family=str(cat.get("family") or ""),
    )
    family = ate_suite or str(cat.get("family") or "")
    live = bool(family) and (bool(ate_suite) or bool(cat.get("live")))
    pk = part_key_for(part)
    model = str(model or part).strip()
    sample = max(1, int(sample_size or 4))
    root = campaign_root(component, part, package, "Version_1", operator=op, base=base)
    created: list[str] = []
    for rel in ("_manifest", "sessions", "workbook"):
        p = root / rel
        if not p.exists():
            p.mkdir(parents=True, exist_ok=True)
            created.append(str(p))
    for dut in range(1, sample + 1):
        for sub in ("screenshots", "graphs"):
            p = root / "Setup" / f"DUT_{dut}" / sub
            if not p.exists():
                p.mkdir(parents=True, exist_ok=True)
                created.append(str(p))

    yaml_path = PARTS_DIR / f"{pk}.yaml"
    wrote_yaml = False
    if base is None and _write_if_missing(yaml_path, _stub_part_yaml(cat, part, package, model, sample)):
        wrote_yaml = True
        created.append(str(yaml_path))
    elif base is not None:
        alt = root / "_manifest" / "part.yaml"
        if _write_if_missing(alt, _stub_part_yaml(cat, part, package, model, sample)):
            wrote_yaml = True
            created.append(str(alt))

    sm = root / "_manifest" / "sheet_map.yaml"
    if _write_if_missing(sm, _tracked_sheet_map(cat, part, package, op, pk, family)):
        created.append(str(sm))
    tc = root / "_manifest" / "test_catalog.yaml"
    if _write_if_missing(tc, _tracked_catalog(pk)):
        created.append(str(tc))
    folders = [folder for _tid, folder in _test_track_rows(pk, family)]
    if not folders:
        folders = ["Setup"]
    created.extend(_ensure_test_folders(root, folders, sample))

    applied = None
    family_error = ""
    ingested: dict[str, Any] | None = None
    if apply and base is None:
        from ate.core.database import set_context

        applied = set_context(
            component=component,
            part=part,
            package=package,
            operator=op,
            version="Version_1",
            model=model,
            part_key=pk,
            sample_size=sample,
        ).identity()
        src_xlsx = resolve_inventory_report(inv)
        have_xlsx = list((root / "workbook").glob("*.xlsx"))
        if src_xlsx is not None and not have_xlsx:
            from ate.core.workbook_import import import_workbook

            ingested = import_workbook(
                source_path=str(src_xlsx),
                dest_name=src_xlsx.name,
                pick=False,
            )
            ingested = _maybe_merge_golden_style(ingested)
    if open_folder and base is None:
        try:
            os.startfile(str(root))
        except OSError:
            pass
    note = (
        "Suite live -- Apply campaign then select tests."
        if live
        else "Folders ready. RUN-IC class has no ATE suite yet (like Level stub)."
    )
    if ingested and ingested.get("workbook"):
        note = (
            f"{note} Workbook copied to campaign workbook/. "
            "Photo/result paste uses sheet_map (FILL_ME until measured). "
            "DEMO does not stamp PASS."
        )
    return {
        "ok": True,
        "root": str(root),
        "created": created,
        "wrote_part_yaml": wrote_yaml,
        "part_key": pk,
        "component": component,
        "operator": op,
        "family": family,
        "sheet_class": str((inv or {}).get("sheet_class") or cat.get("run_ic") or ""),
        "tracking_category": str((inv or {}).get("category") or cat.get("id") or ""),
        "live": live,
        "context": applied,
        "family_error": family_error,
        "workbook": (ingested or {}).get("workbook") or "",
        "sheet_map_action": (ingested or {}).get("sheet_map_action") or "",
        "golden_merged": bool((ingested or {}).get("golden")),
        "note": note,
    }


_LOW_V = frozenset({"vol", "vol_load", "vil"})
_HIGH_V = frozenset({"voh", "voh_load", "vih", "output_voltage"})
_THRESH_V = frozenset({"vih_vil", "vohl", "input_thresholds"})
_I_UA = frozenset(
    {
        "icc",
        "il",
        "iplus",
        "leakage_off",
        "leakage_on",
        "input_leakage",
        "ioff_leakage",
        "off_current",
        "delta_supply_current",
        "supply_current",
        "supply_current_sweep",
        "input_leakage_sweep",
        "ioz",
        "iq",
        "enable_current",
    }
)


def _load_part_raw(part_key: str) -> dict[str, Any]:
    pk = str(part_key or "").strip()
    if not pk:
        return {}
    path = PARTS_DIR / f"{pk}.yaml"
    if not path.is_file():
        return {}
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return raw if isinstance(raw, dict) else {}


def _overlay_part_demo_params(params: dict[str, Any]) -> dict[str, Any]:
    """Fill missing vcc/vccb/current_limit from part yaml. Do not override operator values."""
    out = dict(params)
    raw = _load_part_raw(str(out.get("part") or ""))
    if not raw:
        return out
    if out.get("vcc") in (None, "") and raw.get("vcc", raw.get("vcca")) is not None:
        out["vcc"] = float(raw.get("vcc", raw.get("vcca")))
    if out.get("vccb") in (None, "") and raw.get("vccb") is not None:
        out["vccb"] = float(raw["vccb"])
    if out.get("current_limit_a") in (None, ""):
        ilim = raw.get("current_limit_a", raw.get("current_limit"))
        if ilim is not None:
            out["current_limit_a"] = float(ilim)
    return out


def _order_demo_specs(specs: list[Any]) -> list[Any]:
    """Same board order as START (BUFFER -> G11 -> ORT; unknown modes last)."""
    if len(specs) < 2:
        return list(specs)
    try:
        from ate.core.registry import get as reg_get
        from ate.core.registry import group_by_fixture

        ids = [str(getattr(s, "id", "")) for s in specs]
        if ids and all(reg_get(i) is not None for i in ids):
            batches = group_by_fixture(ids)
            ordered = [s for _mode, group in batches for s in group]
            if ordered:
                return ordered
    except Exception:
        pass
    return list(specs)


def _demo_mock(spec: Any, p: dict[str, Any]) -> dict[str, Any]:
    tid = str(getattr(spec, "id", "") or "")
    need = sorted(getattr(spec, "required_instruments", None) or [])
    vcc = float(p.get("vcc") or 5.0)
    vccb = p.get("vccb")
    vccb_f = float(vccb) if vccb not in (None, "") else None
    ilim = p.get("current_limit_a")
    mock: dict[str, Any] = {"vcc": vcc}
    if vccb_f is not None:
        mock["vccb"] = vccb_f
        if vcc > vccb_f:
            mock["rail_note"] = f"VCCA {vcc} V must be <= VCCB {vccb_f} V"
    if ilim not in (None, ""):
        mock["current_limit_a"] = float(ilim)
    if "DMM" in need:
        if tid in _I_UA:
            mock["i_ua"] = 0.4
        elif tid in _LOW_V:
            mock["dmm_v"] = round(min(0.2, vcc * 0.04), 4)
        elif tid in _THRESH_V:
            mock["dmm_vih"] = round(vcc * 0.7, 4)
            mock["dmm_vil"] = round(vcc * 0.3, 4)
        elif tid in _HIGH_V:
            mock["dmm_v"] = round(vcc * 0.96, 4)
        else:
            mock["dmm_v"] = round(vcc * 0.96, 4)
    if "MSO" in need:
        mock["scope_vpp"] = float(p.get("amp_vpp") or 0.05)
        mock["scope"] = "mock -- no VISA capture"
    if "PSU" in need:
        mock["psu_ch1"] = vcc
        if vccb_f is not None:
            mock["psu_ch2"] = vccb_f
    if "AWG" in need:
        mock["awg_hz"] = float(p.get("freq_hz") or 1000.0)
    if not need:
        mock["note"] = "no instruments required"
    return mock


def build_demo_steps(specs: list[Any], params: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    """Mock instrument plan. No VISA. VOL stays low; dual-rail shows VCCA+VCCB."""
    p = _overlay_part_demo_params(params or {})
    steps: list[dict[str, Any]] = []
    for spec in _order_demo_specs(specs):
        need = sorted(getattr(spec, "required_instruments", None) or [])
        steps.append(
            {
                "test_id": spec.id,
                "label": spec.label,
                "fixture_mode": spec.fixture_mode,
                "instruments": need,
                "lab_sheet": spec.lab_sheet,
                "mock": _demo_mock(spec, p),
            }
        )
    return steps


def run_demo(test_ids: list[str], params: dict[str, Any] | None = None) -> dict[str, Any]:
    """Walk selected tests with mock readings. Writes session JSON + sample files. Not a lab PASS."""
    from ate.core.database import begin_session, end_session, get_context, record_step
    from ate.core.registry import all_tests, get
    from ate.fixture.modes import enabled_tests_for_part

    p = _overlay_part_demo_params(dict(params or {}))
    specs = []
    if test_ids:
        specs = [get(str(i)) for i in test_ids]
        specs = [s for s in specs if s is not None]
    else:
        part_key = str(p.get("part") or "").strip()
        enabled = enabled_tests_for_part(part_key) if part_key else None
        if enabled:
            specs = [s for s in (get(str(i)) for i in enabled) if s is not None]
        else:
            specs = [t for t in all_tests() if t.enabled]
    if not specs:
        raise RuntimeError("No tests to demo -- pick a live family and select tests")
    steps = build_demo_steps(specs, p)
    ctx = get_context()
    folders = [str(s.get("lab_sheet") or s["test_id"]) for s in steps]
    ctx.ensure_tree(folders)
    samples: list[str] = []
    for st in steps:
        folder = str(st.get("lab_sheet") or st["test_id"])
        dest = ctx.graph_dir(folder, 1)
        dest.mkdir(parents=True, exist_ok=True)
        path = dest / "demo_sample.json"
        path.write_text(
            json.dumps({"demo": True, "test_id": st["test_id"], "mock": st["mock"]}, indent=2),
            encoding="utf-8",
        )
        samples.append(str(path))
    p["demo"] = True
    p["test_ids"] = [s.id for s in specs]
    begin_session(
        p,
        instrument_map={"DEMO": "mock", "MSO": "mock", "PSU": "mock", "AWG": "mock", "DMM": "mock"},
        fixture_plan=steps,
    )
    for st in steps:
        from ate.core.specs import any_fail, mock_demo_measurements

        meas = mock_demo_measurements(st["test_id"], part_key=str(ctx.part_key or ""))
        record_step(
            st["test_id"],
            success=not any_fail(meas) if meas else True,
            summary=f"DEMO mock {','.join(st['instruments']) or 'none'}",
            fixture_mode=str(st.get("fixture_mode") or ""),
            measurements=meas or None,
            dut=1,
        )
    session_path = end_session("demo")
    return {
        "ok": True,
        "demo": True,
        "session": str(session_path) if session_path else "",
        "root": str(ctx.root()),
        "steps": steps,
        "samples": samples,
        "note": "Demo only. Mock numbers in sessions/ + DUT graphs/demo_sample.json. Lab xlsx not stamped PASS.",
    }
