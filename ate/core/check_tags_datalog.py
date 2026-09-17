"""Self-check: tags.yaml + TAGS.txt + rolling report.json + archive (A17-T01).

Run: python -m ate.core.check_tags_datalog
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path


def main() -> int:
    errors: list[str] = []
    tmp = Path(tempfile.mkdtemp(prefix="ate_a17_tags_"))
    from ate.core import database as dbmod
    from ate.core import paths as pathmod
    from ate.core.database import begin_session, end_session, record_step, set_context
    from ate.core.datalog import archive_dir, report_path
    from ate.core.datasheet import _merge_specs
    from ate.core.specs import judge_value, load_part_specs
    from ate.core.tags import import_tags, list_boards, load_tags, save_tags, tags_txt_path

    old_root = pathmod.TEST_DB_ROOT
    old_db_root = dbmod.TEST_DB_ROOT
    pathmod.TEST_DB_ROOT = tmp / "#Test_Database"
    dbmod.TEST_DB_ROOT = pathmod.TEST_DB_ROOT
    try:
        pathmod.TEST_DB_ROOT.mkdir(parents=True, exist_ok=True)
        set_context(
            component="OpAmp",
            part="A17Check",
            package="TTSOP8",
            operator="Eugene",
            version="Version_1",
            model="A17CHECK",
            part_key="a17check",
            sample_size=2,
        )
        from ate.core.database import get_context

        ctx = get_context()
        ctx.ensure_tree()

        from openpyxl import Workbook, load_workbook

        xlsx = ctx.workbook_dir() / "A17_Lab_Report.xlsx"
        wb0 = Workbook()
        ws0 = wb0.active
        ws0.title = "Summary"
        ws0["A1"] = "Device"
        ws0["B1"] = "A17Check"
        ws0["A2"] = "Tags"
        wb0.save(xlsx)
        wb0.close()
        ctx.sheet_map_path().write_text(
            "component: OpAmp\npart: A17Check\npackage: TTSOP8\noperator: Eugene\n"
            "version: Version_1\ntests: {}\nidentity:\n  tags: Summary!B2\n",
            encoding="utf-8",
        )

        saved = save_tags(
            ["project:A17-char", "board:G11-REV3"],
            boards=["G11-REV3"],
            ctx=ctx,
        )
        if "project:A17-char" not in saved["tags"]:
            errors.append("save_tags missing project tag")
        if saved.get("excel_status") != "ok":
            errors.append(f"save_tags excel_status={saved.get('excel_status')}")
        else:
            wb_r = load_workbook(xlsx)
            got = wb_r["Summary"]["B2"].value
            wb_r.close()
            if "project:A17-char" not in str(got or ""):
                errors.append("workbook Tags cell missing project token")
        txt = tags_txt_path(ctx)
        if not txt.is_file():
            errors.append("TAGS.txt not written")
        else:
            body = txt.read_text(encoding="utf-8")
            if "board:G11-REV3" not in body:
                errors.append("TAGS.txt missing board token")

        loaded = load_tags(ctx)
        if "project:A17-char" not in loaded["tags"]:
            errors.append("load_tags round-trip failed")

        vocab = list_boards(family="opamp", package="TTSOP8", ctx=ctx)
        if not vocab:
            errors.append("list_boards empty for opamp/TTSOP8")

        pwr_boards = list_boards(family="power", package="SOT23-5", ctx=ctx)
        if not pwr_boards:
            errors.append("list_boards empty for power/SOT23-5")
        if any("G11" in x for x in pwr_boards):
            errors.append("power board vocab must not fall back to OpAmp G11")

        from ate.core.tags import list_label_catalog

        cat = list_label_catalog(family="power", package="SOT23-5", ctx=ctx)
        kinds = {k.get("id") for k in (cat.get("kinds") or []) if isinstance(k, dict)}
        if "board" not in kinds or "board_type" not in kinds:
            errors.append("label kinds must include board and board_type")
        if "LDO" not in ((cat.get("values") or {}).get("board_type") or []):
            errors.append("power board_type vocab must include LDO")

        from ate.core import tags as tagsmod

        old_vocab = tagsmod.LABEL_VOCAB_PATH
        tagsmod.LABEL_VOCAB_PATH = tmp / "label_vocab.yaml"
        try:
            fam_saved = save_tags(
                labels=[{"kind": "tag", "value": "opa-shared"}],
                ctx=ctx,
                remember_scope="family",
            )
            if fam_saved.get("remember_scope") != "family":
                errors.append("save_tags family scope not recorded")
            stored = tagsmod.load_stored_vocab()
            fam_labs = (stored.get("families") or {}).get("opamp") or []
            if not any(x.get("value") == "opa-shared" for x in fam_labs):
                errors.append("family scope must persist in label_vocab.yaml")
            all_saved = save_tags(
                labels=[{"kind": "task", "value": "GBW"}],
                ctx=ctx,
                remember_scope="all",
            )
            if all_saved.get("remember_scope") != "all":
                errors.append("save_tags all scope not recorded")
            cat2 = list_label_catalog(family="opamp", package="TTSOP8", ctx=ctx)
            vals = cat2.get("values") or {}
            if "opa-shared" not in (vals.get("tag") or []):
                errors.append("list_label_catalog must offer class tag opa-shared")
            if "GBW" not in (vals.get("task") or []):
                errors.append("list_label_catalog must offer all-products task GBW")
        finally:
            tagsmod.LABEL_VOCAB_PATH = old_vocab

        labeled = save_tags(
            labels=[{"kind": "board_type", "value": "LDO"}, {"kind": "board", "value": "LDO-SOT23-REV1"}],
            ctx=ctx,
        )
        if "board_type:LDO" not in (labeled.get("tags") or []):
            errors.append("save_tags labels missing board_type:LDO token")
        if "labels" not in labeled:
            errors.append("save_tags must return labels")
        if labeled.get("excel_status") == "ok":
            wb_r2 = load_workbook(xlsx)
            got2 = wb_r2["Summary"]["B2"].value
            wb_r2.close()
            if "board_type:LDO" not in str(got2 or ""):
                errors.append("workbook Tags cell missing board_type after labels save")

        set_context(
            component="OpAmp",
            part="A17Other",
            package="TTSOP8",
            operator="Eugene",
            version="Version_1",
            model="A17OTHER",
            part_key="a17other",
            sample_size=2,
        )
        other = get_context()
        other.ensure_tree()
        save_tags(["project:imported"], boards=[], ctx=other)
        other_root = other.root()

        set_context(
            component="OpAmp",
            part="A17Check",
            package="TTSOP8",
            operator="Eugene",
            version="Version_1",
            model="A17CHECK",
            part_key="a17check",
            sample_size=2,
        )
        ctx = get_context()
        import_tags(other_root, ctx=ctx, merge=True)
        if "project:imported" not in load_tags(ctx)["tags"]:
            errors.append("import_tags failed")

        if judge_value(0.1, -1, 1) != "pass":
            errors.append("judge_value in-window must be pass")
        if judge_value(5, None, 1) != "fail":
            errors.append("judge_value over max must be fail")
        if judge_value(0.5, None, None) != "unspec":
            errors.append("judge_value with no min/max must be unspec")
        merged = _merge_specs(
            [{"id": "VOS_mV", "max": 3.0}],
            [{"id": "VOS_mV", "max": 99, "typ": 0.7}],
        )
        hit = next((x for x in merged if x.get("id") == "VOS_mV"), {})
        if hit.get("max") != 3.0:
            errors.append("datasheet merge must not clobber filled max")
        if hit.get("typ") != 0.7:
            errors.append("datasheet merge must fill empty typ")
        vos = next((x for x in load_part_specs("rs622") if x.get("id") == "VOS_mV"), {})
        if vos.get("max") != 3.0:
            errors.append("rs622 limits yaml missing VOS_mV max 3")

        begin_session({"unit_index": 1, "dut_indices": [1, 2], "run_label": "a17check"})
        record_step(
            "ort",
            success=True,
            summary="ok",
            fixture_mode="G_NEG100",
            dut=1,
            measurements=[{"id": "demo_v", "unit": "V", "min": -1, "max": 1, "value": 0.1}],
        )
        record_step(
            "gbw",
            success=True,
            summary="over max",
            fixture_mode="G11",
            dut=1,
            measurements=[{"id": "demo_hi", "unit": "V", "min": -1, "max": 1, "value": 5}],
        )
        rp = report_path(ctx)
        if not rp.is_file():
            errors.append("report.json missing after record_step")
        else:
            doc = json.loads(rp.read_text(encoding="utf-8"))
            if doc.get("schema") != "ate.datalog.v1":
                errors.append("report schema wrong")
            ident = doc.get("identity") or {}
            if "board_type:LDO" not in (ident.get("tags") or []) and not ident.get("labels"):
                errors.append("session report.json missing labels")
            if not doc.get("steps"):
                errors.append("report steps empty")
            if not any(
                m.get("id") == "demo_v"
                for s in doc.get("sites") or []
                for m in (s.get("measurements") or [])
            ):
                errors.append("site measurement not folded")
            demo_rows = [
                m
                for s in doc.get("steps") or []
                for m in (s.get("measurements") or [])
                if isinstance(m, dict)
            ]
            demo_v = next((m for m in demo_rows if m.get("id") == "demo_v"), {})
            if demo_v.get("result") != "pass":
                errors.append("demo_v must stamp result=pass")
            demo_hi = next((m for m in demo_rows if m.get("id") == "demo_hi"), {})
            if demo_hi.get("result") != "fail":
                errors.append("demo_hi must stamp result=fail")
            gbw_step = next((s for s in doc.get("steps") or [] if s.get("test_id") == "gbw"), {})
            if gbw_step.get("success") is not False:
                errors.append("spec fail must force step success=false")
        md = ctx.sessions_dir() / "datalog.md"
        pdf = ctx.sessions_dir() / "datalog.pdf"
        if not md.is_file():
            errors.append("sessions/datalog.md missing after record_step")
        else:
            body = md.read_text(encoding="utf-8")
            if "PASS" not in body or "FAIL" not in body:
                errors.append("datalog.md must contain PASS and FAIL")
            if "demo_hi" not in body:
                errors.append("datalog.md missing fail param demo_hi")
        if not pdf.is_file() or not pdf.read_bytes().startswith(b"%PDF"):
            errors.append("sessions/datalog.pdf missing or not PDF")
        latest = ctx.root() / "report.pdf"
        if not latest.is_file() or not latest.read_bytes().startswith(b"%PDF"):
            errors.append("Version-root report.pdf missing after record_step")

        end_session("completed")
        arch = archive_dir(ctx)
        archives = list(arch.glob("session_*.json")) if arch.is_dir() else []
        if not archives:
            errors.append("archive missing after end_session")

        # A18: second START with only gbw must keep ort latest + timestamps + records/
        ort_at = None
        doc1 = json.loads(rp.read_text(encoding="utf-8")) if rp.is_file() else {}
        for s in doc1.get("steps") or []:
            if isinstance(s, dict) and s.get("test_id") == "ort":
                ort_at = s.get("at")
        begin_session({"unit_index": 1, "dut_indices": [1], "run_label": "a18merge"})
        record_step(
            "gbw",
            success=True,
            summary="gbw ok",
            fixture_mode="G11",
            dut=1,
            measurements=[{"id": "gbw_mhz", "unit": "MHz", "min": 1, "max": 100, "value": 10}],
        )
        end_session("completed")
        if not rp.is_file():
            errors.append("report.json missing after second START")
        else:
            doc2 = json.loads(rp.read_text(encoding="utf-8"))
            tids = {
                str(s.get("test_id"))
                for s in (doc2.get("steps") or [])
                if isinstance(s, dict)
            }
            if "ort" not in tids:
                errors.append("A18 merge wiped ort after gbw-only START")
            if "gbw" not in tids:
                errors.append("A18 merge missing gbw after second START")
            ort_kept = next(
                (s for s in (doc2.get("steps") or []) if isinstance(s, dict) and s.get("test_id") == "ort"),
                None,
            )
            if ort_at and ort_kept and ort_kept.get("at") != ort_at:
                errors.append("A18 merge changed ort timestamp when gbw-only ran")
            cov = doc2.get("coverage")
            if not isinstance(cov, list):
                errors.append("A18 report missing coverage list")
        ort_records = list((ctx.dut_folder("ort", 1) / "records").glob("ort_*.json"))
        gbw_records = list((ctx.dut_folder("gbw", 1) / "records").glob("gbw_*.json"))
        if not ort_records:
            errors.append("A18 missing ort DUT records/ history file")
        if not gbw_records:
            errors.append("A18 missing gbw DUT records/ history file")

        from ate.core.datalog import delete_run_file, list_runs

        listed = list_runs(scope="campaign", limit=20, ctx=ctx)
        if not listed.get("runs"):
            errors.append("list_runs empty after session")
        else:
            victim = listed["runs"][0]["path"]
            delete_run_file(victim)
            again = list_runs(scope="campaign", limit=20, ctx=ctx)
            if any(r.get("path") == victim for r in again.get("runs") or []):
                errors.append("delete_run_file left the JSON in list_runs")
        try:
            delete_run_file(str(ctx.root() / "workbook" / "nope.json"))
            errors.append("delete_run_file must refuse workbook paths")
        except (PermissionError, ValueError, FileNotFoundError):
            pass
    finally:
        pathmod.TEST_DB_ROOT = old_root
        dbmod.TEST_DB_ROOT = old_db_root
        shutil.rmtree(tmp, ignore_errors=True)

    if errors:
        print("FAIL check_tags_datalog:")
        for e in errors:
            print(f"  - {e}")
        return 1
    print("OK check_tags_datalog")
    return 0


if __name__ == "__main__":
    sys.exit(main())
