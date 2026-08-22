"""Lab-report layout rules (TTSOP golden culture).

Rules (do not drag-resize in Excel):
  - sample_size = 4
  - column width = 16.00 ; row height = 20.4
  - clear ALL borders first (existing cells only — never invent CS* cells)
  - middle+center align everywhere
  - one blank row between major left-intro sections
  - A = meta label, B:L = merged description
  - Test Parameter / Sample Size = 1 row each
  - Test Conditions = wrap-budgeted
  - Test Circuitry = 13 rows fixed
  - Datasheet = 7 rows fixed
  - Test Conclusion = 7 rows fixed
  - after intro finalized: All borders ONLY on A1:L{conclusion_end}
  - Precautions N1:Q4 ; results / output start at N5+
  - after conclusion: 1 gap → indicators (1 row each) → 1 gap → photo boxes
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from ate.reporting.golden_layout import (
    ERROR_TOKENS,
    normalize_breaks,
    rows_for_wrap,
)
from ate.reporting.sheet_layout import (
    HEADER_BLUE,
    HEADER_BLUE_TEMPLATE,
    ORT_LAYOUT,
    YELLOW,
    LayoutSpec,
    _ensure_merge,
    _fill,
    _font_white_bold,
    _safe_unmerge,
    ensure_sample_size,
)

# --- Canonical numbers (user rules) ---
SAMPLE_SIZE = 4
COL_WIDTH = 16.0
ROW_HEIGHT = 20.4
INTRO_END_COL = 12  # L
CIRCUITRY_ROWS = 13
DATASHEET_ROWS = 7
CONCLUSION_ROWS = 7
PHOTO_BODY_ROWS = 10
PHOTO_COLS_PER_BOX = 4  # "max 8 cells" wide band uses 4-col boxes × 2 ch

_NO_BORDER = Border()
_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CH1_FILL = "FFFF00"
_CH2_FILL = "5B9BD5"


@dataclass
class IntroGeometry:
    param_row: int
    sample_row: int
    conditions: tuple[int, int]
    circuitry: tuple[int, int]
    datasheet: tuple[int, int]
    conclusion: tuple[int, int]
    border_end_row: int
    indicator_rows: list[int]
    photo_start_row: int


def clear_borders_existing(ws, max_row: int | None = None, max_col: int | None = None) -> int:
    """Strip borders from already-materialized cells only (no used-range growth).

    Prefer ``ws._cells`` — ``iter_rows`` would create empty cells and re-inflate
    the sheet (CS209-class blowups). Optional max_row/max_col only filter.
    """
    n = 0
    mr = max_row if max_row is not None else 10_000
    mc = max_col if max_col is not None else 10_000
    for (r, c), cell in list(getattr(ws, "_cells", {}).items()):
        if r > mr or c > mc:
            continue
        try:
            if cell.border and any(
                getattr(side, "style", None)
                for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)
            ):
                cell.border = _NO_BORDER
                n += 1
        except AttributeError:
            pass
    return n


def apply_border_box(ws, r0: int, r1: int, c0: int, c1: int) -> None:
    """All-borders on a bounded rectangle (intentional materialization inside box only)."""
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            ws.cell(r, c).border = _THIN


def apply_grid_metrics(ws, *, max_col: int = 32, max_row: int | None = None) -> None:
    for col in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = COL_WIDTH
    last = max_row or min(ws.max_row or 1, 120)
    for r in range(1, last + 1):
        ws.row_dimensions[r].height = ROW_HEIGHT


def _unmerge_left_intro(ws, max_row: int = 80) -> None:
    for mr in list(ws.merged_cells.ranges):
        if mr.min_col <= INTRO_END_COL and mr.max_col <= INTRO_END_COL and mr.min_row <= max_row:
            _safe_unmerge(ws, str(mr))


def _unmerge_photo_band(ws, start_row: int, end_row: int) -> None:
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= start_row and mr.max_row <= end_row and mr.min_col <= 32:
            _safe_unmerge(ws, str(mr))


def _read_label_value(ws, label: str, scan_to: int = 80) -> Any:
    key = label.strip().lower()
    for r in range(1, scan_to + 1):
        lab = ws.cell(r, 1).value
        if lab and str(lab).strip().lower() == key:
            return ws.cell(r, 2).value
    return None


def _style_meta_label(cell) -> None:
    cell.fill = _fill(HEADER_BLUE_TEMPLATE)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.alignment = _CENTER


def _write_intro_block(ws, label: str, value: Any, r0: int, r1: int) -> None:
    for r in range(r0, r1 + 1):
        for c in range(1, INTRO_END_COL + 1):
            cell = ws.cell(r, c)
            try:
                if not (r == r0 and c in (1, 2)):
                    cell.value = None
            except AttributeError:
                pass
            cell.alignment = _CENTER
    lab = ws.cell(r0, 1)
    lab.value = label
    _style_meta_label(lab)
    val = ws.cell(r0, 2)
    text = normalize_breaks(value)
    if text.strip().upper() in ERROR_TOKENS:
        text = ""
    val.value = text if text else None
    val.alignment = _CENTER
    _ensure_merge(ws, f"A{r0}:A{r1}")
    _ensure_merge(ws, f"B{r0}:{get_column_letter(INTRO_END_COL)}{r1}")


def compute_intro_geometry(
    conditions_text: Any,
    *,
    indicator_count: int = 2,
) -> IntroGeometry:
    """Stack left intro with 1-row gaps between major sections."""
    cond = rows_for_wrap(
        conditions_text,
        merge_cols=11,
        col_width=COL_WIDTH,
        min_rows=4,
        max_rows=18,
        pad_rows=1,
    )
    r = 1
    param_row = r
    r += 1
    sample_row = r
    r += 1
    cond0 = r
    cond1 = cond0 + cond.rows - 1
    r = cond1 + 2  # gap
    circ0 = r
    circ1 = circ0 + CIRCUITRY_ROWS - 1
    r = circ1 + 2
    ds0 = r
    ds1 = ds0 + DATASHEET_ROWS - 1
    r = ds1 + 2
    conc0 = r
    conc1 = conc0 + CONCLUSION_ROWS - 1
    border_end = conc1

    # gap then indicators
    ind_start = conc1 + 2
    indicator_rows = list(range(ind_start, ind_start + max(1, indicator_count)))
    photo_start = indicator_rows[-1] + 2
    return IntroGeometry(
        param_row=param_row,
        sample_row=sample_row,
        conditions=(cond0, cond1),
        circuitry=(circ0, circ1),
        datasheet=(ds0, ds1),
        conclusion=(conc0, conc1),
        border_end_row=border_end,
        indicator_rows=indicator_rows,
        photo_start_row=photo_start,
    )


def ensure_precautions_n1q4(ws) -> None:
    """Precautions locked to N1:Q4 so results can start at N5."""
    for mr in list(ws.merged_cells.ranges):
        # Unmerge anything that ate into row 5+ in the precaution columns
        if mr.min_col >= 14 and mr.min_col <= 18 and mr.min_row <= 5:
            _safe_unmerge(ws, str(mr))
    ws.cell(1, 14).value = ws.cell(1, 14).value or "Precautions"
    _style_meta_label(ws.cell(1, 14))
    # Wipe stub values in the precaution body before merge
    for r in range(2, 6):
        for c in range(14, 19):
            cell = ws.cell(r, c)
            try:
                if cell.value is not None and str(cell.value).strip() in {"1", "#1"}:
                    cell.value = None
            except AttributeError:
                pass
    _ensure_merge(ws, "N1:Q1")
    _ensure_merge(ws, "N2:Q4")
    ws.cell(2, 14).alignment = _CENTER
    apply_border_box(ws, 1, 4, 14, 17)
    # Ensure N5 is not part of precaution merge / stub values
    for c in range(14, 18):
        cell = ws.cell(5, c)
        try:
            if cell.value is not None and str(cell.value).strip() in {"1", "#1"}:
                cell.value = None
            cell.border = _NO_BORDER
        except AttributeError:
            pass


def place_indicators(ws, rows: list[int], items: list[tuple[str, str]] | None = None) -> None:
    """One indicator per row (CH1/CH2 by default)."""
    items = items or [("CH1: IN+", _CH1_FILL), ("CH2: VOUT", _CH2_FILL)]
    # Optional section tag on the gap row above first indicator
    tag_row = rows[0] - 1
    if tag_row >= 1 and (ws.cell(tag_row, 1).value in (None, "")):
        # leave blank gap as required; do not put "Indicators" into the gap
        pass
    for i, row in enumerate(rows):
        label, fill = items[i] if i < len(items) else (f"Indicator {i+1}", YELLOW)
        cell = ws.cell(row, 1)
        cell.value = label
        cell.fill = _fill(fill)
        cell.font = Font(bold=True, size=11)
        cell.alignment = _CENTER
        apply_border_box(ws, row, row, 1, 1)


def apply_photo_grid_at(
    ws,
    start_row: int,
    *,
    title: str = "Overload Recovery Time",
    sample_size: int = SAMPLE_SIZE,
    body_rows: int = PHOTO_BODY_ROWS,
    dual_polarity: bool = True,
) -> dict[str, str]:
    """Photo boxes after a 1-row gap from indicators. ORT uses POS then NEG bands."""
    yellow = _fill(YELLOW)
    blue = _fill(HEADER_BLUE)
    box_starts = (1, 5, 9, 13, 17, 21, 25, 29)[: sample_size * 2]
    unit_starts = (1, 9, 17, 25)[:sample_size]
    anchors: dict[str, str] = {}

    def _one_band(header_row: int, band_title: str, key_prefix: str) -> int:
        channel_row = header_row + 1
        body0 = header_row + 2
        body1 = body0 + body_rows - 1
        for i, start in enumerate(unit_starts, start=1):
            end = start + 7
            _ensure_merge(
                ws,
                f"{get_column_letter(start)}{header_row}:{get_column_letter(end)}{header_row}",
            )
            cell = ws.cell(header_row, start)
            cell.value = f"{band_title} #{i}"
            cell.fill = blue
            cell.font = _font_white_bold()
            cell.alignment = _CENTER
        channels = ("Channel A", "Channel B")
        for box_i, start_col in enumerate(box_starts):
            end_col = start_col + PHOTO_COLS_PER_BOX - 1
            unit = box_i // 2 + 1
            ch_key = "chA" if box_i % 2 == 0 else "chB"
            _ensure_merge(
                ws,
                f"{get_column_letter(start_col)}{channel_row}:{get_column_letter(end_col)}{channel_row}",
            )
            c = ws.cell(channel_row, start_col)
            c.value = channels[box_i % 2]
            c.fill = yellow
            c.font = Font(bold=True)
            c.alignment = _CENTER
            apply_border_box(ws, channel_row, channel_row, start_col, end_col)
            for r in range(body0, body1 + 1):
                for col in range(start_col, end_col + 1):
                    cell = ws.cell(r, col)
                    try:
                        if cell.value is not None and str(cell.value).strip().upper() in ERROR_TOKENS:
                            cell.value = None
                    except AttributeError:
                        pass
                    cell.alignment = _CENTER
            _ensure_merge(
                ws,
                f"{get_column_letter(start_col)}{body0}:{get_column_letter(end_col)}{body1}",
            )
            apply_border_box(ws, body0, body1, start_col, end_col)
            anchors[f"{key_prefix}_u{unit}_{ch_key}"] = f"{get_column_letter(start_col)}{body0}"
        for r in range(header_row, body1 + 1):
            ws.row_dimensions[r].height = ROW_HEIGHT
        return body1

    if dual_polarity:
        end_pos = _one_band(start_row, f"Positive {title}", "pos")
        neg_start = end_pos + 2  # one-row gap between POS/NEG
        _one_band(neg_start, f"Negative {title}", "neg")
    else:
        _one_band(start_row, title, "u")
    return anchors


def apply_rules_ort_sheet(ws, *, sample_size: int = SAMPLE_SIZE) -> dict:
    """Apply TTSOP golden rules to the ORT sheet without rewriting drawings."""
    # 1) Clear borders on existing footprint only (cap inflated sheets)
    cleared = clear_borders_existing(ws)

    # 2) Capture intro values before unmerge
    param = _read_label_value(ws, "Test Parameter") or "Overload Recovery Time"
    sample = _read_label_value(ws, "Sample Size") or sample_size
    conditions = _read_label_value(ws, "Test Conditions")
    circuitry = _read_label_value(ws, "Test Circuitry")
    datasheet = _read_label_value(ws, "Datasheet")
    conclusion = _read_label_value(ws, "Test Conclusion")
    if isinstance(conditions, str):
        conditions = (
            conditions.replace("IN- offset = ±100V", "IN- offset = ±100 mV")
            .replace("IN- offset = ±100 V", "IN- offset = ±100 mV")
            .replace("IN- offset = +/-100V", "IN- offset = ±100 mV")
        )
    for bad in (circuitry, datasheet):
        if bad is not None and str(bad).strip().upper() in ERROR_TOKENS:
            if bad is circuitry:
                circuitry = None
            else:
                datasheet = None

    geo = compute_intro_geometry(conditions, indicator_count=2)

    # 3) Clear old left merges + old photo band (keep right N5+ results intact)
    old_photo_guess = 40
    _unmerge_left_intro(ws, max_row=max(90, geo.photo_start_row + 40))
    _unmerge_photo_band(ws, old_photo_guess, max(ws.max_row or 68, geo.photo_start_row + 40))

    # Clear stale left labels below new geometry that might linger
    for r in range(1, max(90, geo.photo_start_row + 30)):
        for c in range(1, INTRO_END_COL + 1):
            cell = ws.cell(r, c)
            try:
                cell.value = None
            except AttributeError:
                pass

    # 4) Write stacked intro with gaps
    _write_intro_block(ws, "Test Parameter", param, geo.param_row, geo.param_row)
    _write_intro_block(ws, "Sample Size", sample_size, geo.sample_row, geo.sample_row)
    _write_intro_block(ws, "Test Conditions", conditions, *geo.conditions)
    _write_intro_block(ws, "Test Circuitry", circuitry, *geo.circuitry)
    _write_intro_block(ws, "Datasheet", datasheet, *geo.datasheet)
    _write_intro_block(ws, "Test Conclusion", conclusion, *geo.conclusion)

    # 5) Borders ONLY on A1:L{conclusion_end}
    apply_border_box(ws, 1, geo.border_end_row, 1, INTRO_END_COL)

    # 6) Precautions + N5 rule
    ensure_precautions_n1q4(ws)
    ensure_sample_size(ws, sample_size)

    # 7) Indicators then photo grid
    place_indicators(ws, geo.indicator_rows)
    anchors = apply_photo_grid_at(
        ws,
        geo.photo_start_row,
        title="Overload Recovery Time",
        sample_size=sample_size,
        dual_polarity=True,
    )

    # 8) Grid metrics
    photo_end = geo.photo_start_row + (2 + PHOTO_BODY_ROWS) * 2 + 2
    apply_grid_metrics(ws, max_col=32, max_row=photo_end)

    return {
        "mode": "rules_ort",
        "borders_cleared": cleared,
        "geometry": {
            "conditions": geo.conditions,
            "circuitry": geo.circuitry,
            "datasheet": geo.datasheet,
            "conclusion": geo.conclusion,
            "border_end_row": geo.border_end_row,
            "indicator_rows": geo.indicator_rows,
            "photo_start_row": geo.photo_start_row,
        },
        "photo_anchors": anchors,
        "sample_size": sample_size,
        "col_width": COL_WIDTH,
        "row_height": ROW_HEIGHT,
    }


def patch_layout_spec_defaults(spec: LayoutSpec | None = None) -> LayoutSpec:
    """Align shared LayoutSpec numeric defaults with golden rules."""
    spec = spec or ORT_LAYOUT
    spec.sample_size = SAMPLE_SIZE
    spec.standard_col_width = COL_WIDTH
    spec.label_col_width = COL_WIDTH
    spec.photo_row_height = ROW_HEIGHT
    return spec
