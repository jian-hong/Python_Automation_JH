"""Apply ORT 16-photo-box layout + sample_size=4 across lab report sheets."""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from openpyxl import load_workbook

from ate.reporting.sheet_layout import (
    apply_full_ort_layout,
    light_standardize_sheet,
)

LAB = Path(r"C:\Users\OoiJianHong\Downloads\RS622XK_Lab_Report_TTSOP.xlsx")
BACKUP = Path(r"C:\Users\OoiJianHong\Downloads\RS622XK_Lab_Report_TTSOP.backup.xlsx")
DB_LAB = Path(
    r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1\workbook\RS622XK_Lab_Report_TTSOP.xlsx"
)
PRIMARY = DB_LAB if DB_LAB.is_file() else LAB
MANIFEST = Path(
    r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1\_manifest\sheet_map.yaml"
)
DB_ROOT = Path(r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1")

TEST_SHEETS = [
    "Slew Rate",
    "GBW",
    "SettlingTime",
    "SSSR",
    "LSSR",
    "ORT",
    "PSRR",
    "CMRR",
    "AOL",
    "EMIRR",
    "NoPhaseReversal",
    "VOS",
    "PowerOnTime",
    "VOL",
    "Noise",
]


def main() -> int:
    src = PRIMARY if PRIMARY.is_file() else (LAB if LAB.is_file() else DB_LAB)
    if not src.is_file():
        print(f"Missing lab report: {src}")
        return 1

    if LAB.is_file() and not BACKUP.exists():
        shutil.copy2(LAB, BACKUP)
        print(f"Backup: {BACKUP}")

    for folder in (
        "ORT", "SlewRate", "GBW", "SettlingTime", "SSSR", "LSSR", "VOS",
        "PowerOnTime", "VOL", "EMIRR", "PSRR", "CMRR", "AOL", "NoPhaseReversal", "Noise",
    ):
        (DB_ROOT / folder).mkdir(parents=True, exist_ok=True)

    wb = load_workbook(src)

    for name in TEST_SHEETS:
        if name not in wb.sheetnames:
            continue
        if name == "ORT":
            continue
        light_standardize_sheet(wb[name], sample_size=4)
        print(f"  light: {name}")

    anchors = {}
    if "ORT" in wb.sheetnames:
        anchors = apply_full_ort_layout(wb["ORT"])
        print("ORT 16 photo boxes:")
        print(json.dumps(anchors, indent=2))

    DB_LAB.parent.mkdir(parents=True, exist_ok=True)
    alt = DB_ROOT / "workbook" / "RS622XK_Lab_Report_TTSOP_layout16.xlsx"
    try:
        wb.save(DB_LAB)
        out = DB_LAB
        print(f"Saved: {out}")
        if LAB.is_file():
            try:
                shutil.copy2(out, LAB)
                print(f"Synced Downloads: {LAB}")
            except Exception as exc:
                print(f"Downloads sync skipped: {exc}")
    except PermissionError:
        try:
            wb.save(LAB)
            out = LAB
            print(f"DB locked — saved: {out}")
        except PermissionError:
            wb.save(alt)
            out = alt
            print(f"Both locked — saved unlocked copy: {out}")
            print("Close Excel, then replace the workbook with this file.")
    wb.close()

    if MANIFEST.is_file() and anchors:
        try:
            import yaml

            data = yaml.safe_load(MANIFEST.read_text(encoding="utf-8")) or {}
            ort = data.setdefault("tests", {}).setdefault("ORT", {})
            photos: dict = {}
            for key, cell in anchors.items():
                parts = key.split("_")
                pol = parts[0].upper()
                unit = parts[1].replace("u", "")
                ch = parts[2].upper()
                dut = f"DUT_{unit}"
                photos.setdefault(dut, {})[f"{pol}_{ch}"] = cell
            paste = ort.setdefault("paste", {})
            paste["photos"] = photos
            paste["results"] = {
                "pos_cha": ["R13", "S13", "T13", "U13"],
                "pos_chb": ["AA13", "AB13", "AC13", "AD13"],
                "neg_cha": ["R20", "S20", "T20", "U20"],
                "neg_chb": ["AA20", "AB20", "AC20", "AD20"],
                "average": "Q7",
            }
            paste["setup_text"] = "B3"
            paste["conclusion"] = "B34"
            ort["dut_iterations"] = 4
            data["sample_size"] = 4
            data["layout_rules"] = {
                "results_zone": "right-top",
                "precautions_zone": "right-top",
                "schematic_zone": "left-mid",
                "screenshot_zone": "bottom-16-boxes",
                "merge_center": True,
                "equal_cell_grid": True,
                "photo_box_px": {"width": 420, "height": 240},
                "sample_size": 4,
            }
            MANIFEST.write_text(
                yaml.safe_dump(data, sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            print(f"Updated manifest: {MANIFEST}")
        except Exception as exc:
            print(f"Manifest update skipped: {exc}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
