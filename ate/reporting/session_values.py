"""Fill campaign workbook number cells from living report.json.

Path B (excel_plots + workbook_policy one_per_version_overwrite): create or
overwrite the one Version xlsx. Never a second orphan book / _filled.xlsx.
OpAmp / mapped families: sheet_map paste.values (measurement id -> cell,
DUT list, or CHA/CHB grid). Skips FILL_ME. Does not run OpAmp golden on Path B.
"""
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from ate.reporting.session_paste import _is_real_cell

from openpyxl.cell.cell import MergedCell

_LOG = logging.getLogger("ate.session_values")
_CELL_RE = re.compile(r"^[A-Z]{1,3}\d{1,5}$")


def _cell_token(raw: Any) -> str:
    s = str(raw or "").split("#", 1)[0].strip().upper()
    return s if _CELL_RE.match(s) else ""


def _sheet_and_cell(token: str, default_sheet: str) -> tuple[str, str]:
    if "!" in token:
        sh, cell = token.split("!", 1)
        return sh.strip() or default_sheet, cell.strip()
    return default_sheet, token


def _dut_index(step: dict[str, Any]) -> int:
    try:
        return max(1, int(step.get("dut") or 1))
    except (TypeError, ValueError):
        return 1


def _channel_key(step: dict[str, Any]) -> str:
    raw = str(step.get("channel") or "CHA").strip().upper()
    if raw in ("CHB", "B", "2"):
        return "CHB"
    return "CHA"


def _expand_cells(raw: Any, dut: int, channel: str) -> list[str]:
    """One cell for this DUT/channel. List is DUT-indexed; dict is CHA/CHB."""
    if raw is None:
        return []
    if isinstance(raw, dict):
        pick = (
            raw.get(channel)
            or raw.get(channel.lower())
            or raw.get("CHA")
            or raw.get("cha")
        )
        return _expand_cells(pick, dut, channel)
    if isinstance(raw, (list, tuple)):
        i = dut - 1
        if i < 0 or i >= len(raw):
            return []
        item = raw[i]
        if not _is_real_cell(item):
            return []
        tok = _cell_token(item)
        return [tok] if tok else []
    if not _is_real_cell(raw):
        return []
    tok = _cell_token(raw)
    return [tok] if tok else []


def _fold_name(raw: str) -> set[str]:
    """voh_load -> {vohload, voh}; Supply_Current -> {supplycurrent}."""
    n = re.sub(r"[^a-z0-9]+", "", str(raw or "").lower())
    if not n:
        return set()
    out = {n}
    for suf in ("load", "sweep"):
        if n.endswith(suf) and len(n) > len(suf):
            out.add(n[: -len(suf)])
    return out


def _entry_for_step(tests: dict[str, Any], step: dict[str, Any]) -> tuple[str, dict[str, Any]] | None:
    tid = str(step.get("test_id") or "").strip()
    if not tid:
        return None
    want = [tid, str(step.get("lab_sheet") or ""), str(step.get("fixture_mode") or "")]
    want_fold = set()
    for a in want:
        want_fold |= _fold_name(a)
    for key, entry in tests.items():
        if not isinstance(entry, dict):
            continue
        names = [str(key), str(entry.get("folder") or ""), str(entry.get("excel_sheet") or "")]
        name_fold = set()
        for b in names:
            name_fold |= _fold_name(b)
        if want_fold & name_fold:
            return str(key), entry
        for a in want:
            if not a:
                continue
            for b in names:
                if a.lower() == b.lower() or a.lower().replace("_", "") == b.lower().replace("_", ""):
                    return str(key), entry
    return None


def fill_workbook_from_report(
    *,
    report: dict[str, Any] | None = None,
    workbook_path: Path | None = None,
    ctx=None,
) -> dict[str, Any]:
    from ate.core.database import get_context
    from ate.core.datalog import load_report

    c = ctx or get_context()
    doc = report if isinstance(report, dict) else load_report(ctx=c)
    path = Path(workbook_path) if workbook_path else c.lab_report_path()

    try:
        from ate.tests.logic.product_model import has_product_model, load_product_model
        from ate.tests.logic.excel_lock import (
            OrphanWorkbook,
            uses_excel_lock,
            write_path_b_workbook,
        )

        key = str(getattr(c, "part_key", "") or "").strip().lower()
        if key and has_product_model(key):
            model = load_product_model(key)
            if uses_excel_lock(model):
                try:
                    return write_path_b_workbook(ctx=c, model=model, report=doc)
                except OrphanWorkbook as exc:
                    return {
                        "filled": 0,
                        "skipped": 0,
                        "status": "orphan",
                        "excel": str(path),
                        "error": str(exc),
                    }
                except Exception as exc:
                    _LOG.warning("Path B workbook: %s", exc)
                    return {
                        "filled": 0,
                        "skipped": 0,
                        "status": "error",
                        "excel": str(path),
                        "error": str(exc),
                    }
    except Exception:
        pass

    sm = c.load_sheet_map() if hasattr(c, "load_sheet_map") else {}
    tests = sm.get("tests") if isinstance(sm, dict) else {}
    if not isinstance(tests, dict):
        return {"filled": 0, "skipped": 0, "status": "no_sheet_map"}
    if not path.is_file():
        return {"filled": 0, "skipped": 0, "status": "no_workbook", "excel": str(path)}

    writes: list[tuple[str, str, Any]] = []
    skipped = 0
    for step in doc.get("steps") or []:
        if not isinstance(step, dict):
            continue
        hit = _entry_for_step(tests, step)
        if not hit:
            skipped += 1
            continue
        _key, entry = hit
        paste = entry.get("paste") if isinstance(entry.get("paste"), dict) else {}
        sheet = str(entry.get("excel_sheet") or _key)
        values = paste.get("values") if isinstance(paste.get("values"), dict) else {}
        meas = step.get("measurements") if isinstance(step.get("measurements"), list) else []
        if not values or not meas:
            skipped += 1
            continue
        dut = _dut_index(step)
        channel = _channel_key(step)
        result_written = False
        for m in meas:
            if not isinstance(m, dict):
                continue
            mid = str(m.get("id") or "")
            if mid and mid in values:
                for cell in _expand_cells(values[mid], dut, channel):
                    writes.append((*_sheet_and_cell(cell, sheet), m.get("value")))
            if result_written:
                continue
            for cell in _expand_cells(values.get("result"), dut, channel):
                if m.get("result"):
                    writes.append((*_sheet_and_cell(cell, sheet), str(m.get("result")).upper()))
                    result_written = True

    if not writes:
        return {"filled": 0, "skipped": skipped, "status": "nothing_mapped", "excel": str(path)}

    from openpyxl import load_workbook

    wb = None
    saved = path
    status = "ok"
    try:
        wb = load_workbook(path)
        n = 0
        for sh, cell, val in writes:
            if sh not in wb.sheetnames or not cell:
                skipped += 1
                continue
            target = wb[sh][cell]
            if isinstance(target, MergedCell):
                skipped += 1
                continue
            target.value = val
            n += 1
        try:
            wb.save(path)
        except PermissionError:
            alt = path.with_name(path.stem + "_filled.xlsx")
            wb.save(alt)
            saved = alt
            status = "locked"
        return {
            "filled": n,
            "skipped": skipped,
            "status": status,
            "excel": str(saved),
        }
    except PermissionError as exc:
        return {"filled": 0, "skipped": skipped, "status": "locked", "excel": str(path), "error": str(exc)}
    except Exception as exc:
        _LOG.warning("fill workbook: %s", exc)
        return {"filled": 0, "skipped": skipped, "status": "error", "excel": str(path), "error": str(exc)}
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
