"""Apply golden ORT layout culture to every test sheet in the lab report."""
from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from openpyxl import load_workbook

from ate.reporting.golden_layout import TEST_SHEETS, apply_golden_sheet, rows_for_wrap

LAB = Path(r"C:\Users\OoiJianHong\Downloads\RS622XK_Lab_Report_TTSOP.xlsx")
BACKUP = Path(r"C:\Users\OoiJianHong\Downloads\RS622XK_Lab_Report_TTSOP.backup.xlsx")
DB_LAB = Path(
    r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1\workbook\RS622XK_Lab_Report_TTSOP.xlsx"
)
OUT_ALT = Path(
    r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1\workbook\RS622XK_Lab_Report_TTSOP_golden.xlsx"
)
MANIFEST = Path(
    r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1\_manifest\sheet_map.yaml"
)
REPORT_JSON = Path(
    r"C:\Users\OoiJianHong\#Test_Database\OpAmp\RS622\TTSOP8\Version_1\_manifest\golden_layout_report.json"
)
SAMPLE_SIZE = 4


def _pick_src() -> Path:
    # Prefer pristine backup so re-runs don't double-insert wrap rows
    for p in (BACKUP, LAB, DB_LAB, OUT_ALT):
        if p.is_file():
            return p
    raise FileNotFoundError("No lab report workbook found")


def main() -> int:
    src = _pick_src()
    print(f"Source: {src}")

    if LAB.is_file() and not BACKUP.exists():
        shutil.copy2(LAB, BACKUP)
        print(f"Backup: {BACKUP}")

    wb = load_workbook(src)
    report: dict = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "source": str(src),
        "sample_size": SAMPLE_SIZE,
        "wrap_examples": {},
        "sheets": {},
    }

    # Document wrap math for a few known long texts
    if "ORT" in wb.sheetnames:
        cond = wb["ORT"]["B3"].value
        budget = rows_for_wrap(cond, merge_cols=11, min_rows=4, max_rows=18)
        report["wrap_examples"]["ORT_B3"] = {
            "chars_per_line": budget.chars_per_line,
            "lines": budget.lines,
            "rows": budget.rows,
        }

    for name in TEST_SHEETS:
        if name not in wb.sheetnames:
            print(f"  skip missing: {name}")
            continue
        info = apply_golden_sheet(wb[name], name, sample_size=SAMPLE_SIZE)
        report["sheets"][name] = info
        print(f"  {name}: mode={info.get('mode')} photos@{info.get('photo_start_row', 'ORT')}")

    # Summary sample size reminder
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        for row in range(1, 40):
            if ws.cell(row, 1).value and "sample" in str(ws.cell(row, 1).value).lower():
                ws.cell(row, 2).value = SAMPLE_SIZE

    REPORT_JSON.parent.mkdir(parents=True, exist_ok=True)
    REPORT_JSON.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    print(f"Report: {REPORT_JSON}")

    saved = None
    for dest in (DB_LAB, OUT_ALT, LAB):
        try:
            dest.parent.mkdir(parents=True, exist_ok=True)
            wb.save(dest)
            saved = dest
            print(f"Saved: {dest}")
            break
        except PermissionError:
            print(f"Locked: {dest}")
            continue
    wb.close()

    if saved is None:
        print("FAILED: all save targets locked. Close Excel and re-run.")
        return 2

    # Sync other copies when possible
    for dest in (DB_LAB, LAB, OUT_ALT):
        if dest == saved:
            continue
        try:
            shutil.copy2(saved, dest)
            print(f"Synced: {dest}")
        except Exception as exc:
            print(f"Sync skip {dest.name}: {exc}")

    # Update manifest sample_size + layout_rules
    if MANIFEST.is_file():
        try:
            import yaml

            data = yaml.safe_load(MANIFEST.read_text(encoding="utf-8")) or {}
            data["sample_size"] = SAMPLE_SIZE
            data["layout_rules"] = {
                "results_zone": "right-top",
                "precautions_zone": "right-top N1:Q4",
                "schematic_zone": "left-mid (wrap-merge)",
                "screenshot_zone": "bottom photo grid (4 DUT x ChA/ChB)",
                "merge_center": True,
                "wrap_text_sizing": True,
                "equal_cell_grid": True,
                "photo_box": {"cols": 4, "rows": 10, "px": {"width": 420, "height": 240}},
                "sample_size": SAMPLE_SIZE,
            }
            # Attach photo anchors per sheet from report
            for sname, info in report["sheets"].items():
                # Map excel sheet name to manifest test keys loosely
                key = sname.replace(" ", "")
                tests = data.setdefault("tests", {})
                # Find matching test entry by excel_sheet
                target = None
                for tkey, tval in tests.items():
                    if isinstance(tval, dict) and tval.get("excel_sheet") in {sname, key, sname.replace(" Rate", "Rate")}:
                        target = tval
                        break
                if target is None:
                    continue
                paste = target.setdefault("paste", {})
                if info.get("photo_anchors"):
                    paste["photos"] = info["photo_anchors"]
                target["dut_iterations"] = SAMPLE_SIZE
            MANIFEST.write_text(
                yaml.safe_dump(data, sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            print(f"Updated: {MANIFEST}")
        except Exception as exc:
            print(f"Manifest update skipped: {exc}")

    print(f"DONE saved={saved}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
