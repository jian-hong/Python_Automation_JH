"""Path B Excel lock: golden_auto Version book vs ultimate_manual jot.

golden_auto = chosen campaign Version
`#Test_Database/.../{Version_N}/workbook/` -- Continue/START overwrite-in-place.
JSON->Excel from card-backed field ids only. Never invent columns.

ultimate_manual = separate all-test jot workbook -- NEVER the auto target.
workbook_policy.golden_auto: one_per_version_overwrite
workbook_policy.ultimate_manual: never_auto_write

A second golden file in Version workbook/ is an orphan FAIL.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from ate.tests.logic.product_model import ProductModel, icc_pins

WORKBOOK_POLICY = "one_per_version_overwrite"
ULTIMATE_POLICY = "never_auto_write"
PRETTY_POLICY = ULTIMATE_POLICY

# Filename / folder tokens that mark the jot / pretty book. Do not write auto runs there.
_ULTIMATE_RX = re.compile(
    r"(ultimate(_?manual)?|(^|[^a-z0-9])(jot|pretty)([^a-z0-9]|$)|all[-_]?test)",
    re.I,
)

# Card-backed series ids. Do not add noise / GBW / invented plots.
ALLOWED_SERIES = frozenset(
    {
        "vih_vs_vcc",
        "vil_vs_vcc",
        "vtplus_vs_vcc",
        "vtminus_vs_vcc",
        "dvt_vs_vcc",
        "icc_vs_vcc",
        "delta_icc_vs_vcc",
        "ii_vs_vcc",
        "voh_at_ioh",
        "vol_at_iol",
        "ioz_vs_vcc",
    }
)

SERIES_TEST = {
    "vih_vs_vcc": "input_threshold",
    "vil_vs_vcc": "input_threshold",
    "vtplus_vs_vcc": "input_threshold",
    "vtminus_vs_vcc": "input_threshold",
    "dvt_vs_vcc": "input_threshold",
    "icc_vs_vcc": "icc",
    "delta_icc_vs_vcc": "delta_icc",
    "ii_vs_vcc": "ii",
    "voh_at_ioh": "voh",
    "vol_at_iol": "vol",
    "ioz_vs_vcc": "ioz",
}

# TestSpec.lab_sheet from logic_dc._register. Not invented campaign cells.
_SHEET = {
    "input_threshold": "VTH",
    "vth": "VTH",
    "icc": "Icc",
    "delta_icc": "DeltaICC",
    "ii": "II",
    "voh": "VOH",
    "vol": "VOL",
    "ioz": "IOZ",
}

_PATH_B_DC = frozenset(
    {"input_threshold", "vth", "icc", "delta_icc", "ii", "voh", "vol", "ioz"}
)

# Regex: detect plottable series from known runner headers only.
_HEADER_RX = (
    (re.compile(r"^VIH(_min_V)?$", re.I), "vih_vs_vcc"),
    (re.compile(r"^VIL(_max_V)?$", re.I), "vil_vs_vcc"),
    (re.compile(r"^(VT\+|VTPLUS(_V)?)$", re.I), "vtplus_vs_vcc"),
    (re.compile(r"^(VT-|VTMINUS(_V)?)$", re.I), "vtminus_vs_vcc"),
    (re.compile(r"^(HYSTERESIS_V|HYST|DVT|DVT_V)$", re.I), "dvt_vs_vcc"),
    (re.compile(r"^ICC_uA$", re.I), "icc_vs_vcc"),
    (re.compile(r"^DELTA_ICC_uA$", re.I), "delta_icc_vs_vcc"),
    (re.compile(r"^II_uA$", re.I), "ii_vs_vcc"),
    (re.compile(r"^(IOH_A|Measured|Spec_min|VOH)$", re.I), "voh_at_ioh"),
    (re.compile(r"^(IOL_A|Spec_max|VOL)$", re.I), "vol_at_iol"),
    (re.compile(r"^IOZ_uA$", re.I), "ioz_vs_vcc"),
)

_Y_FOR_SERIES = {
    "vih_vs_vcc": ("VIH", "VIH_min_V"),
    "vil_vs_vcc": ("VIL", "VIL_max_V"),
    "vtplus_vs_vcc": ("VT+",),
    "vtminus_vs_vcc": ("VT-",),
    "dvt_vs_vcc": ("HYSTERESIS_V",),
    "icc_vs_vcc": ("ICC_uA",),
    "delta_icc_vs_vcc": ("ICC_uA", "DELTA_ICC_uA"),
    "ii_vs_vcc": ("II_uA",),
    "voh_at_ioh": ("Measured", "VOH", "Spec_min"),
    "vol_at_iol": ("Measured", "VOL", "Spec_max"),
    "ioz_vs_vcc": ("IOZ_uA",),
}

_X_FOR_SERIES = {
    "voh_at_ioh": "IOH_A",
    "vol_at_iol": "IOL_A",
}


class OrphanWorkbook(RuntimeError):
    """A run would create or leave a second golden xlsx under workbook/."""


class UltimateWorkbook(RuntimeError):
    """Auto write targeted the ultimate_manual jot book."""


def workbook_policy_map(model: ProductModel | None) -> dict[str, str]:
    """auto = golden_auto overwrite Version; pretty = never auto (ultimate_manual)."""
    if model is None:
        return {}
    out: dict[str, str] = {}
    raw = getattr(model, "workbook_policy", None)
    if isinstance(raw, dict):
        out = {
            str(k).strip(): str(v or "").strip()
            for k, v in raw.items()
            if str(k).strip()
        }
    elif isinstance(raw, str) and raw.strip():
        out = {"golden_auto": raw.strip()}
    else:
        plots = getattr(model, "excel_plots", None)
        if isinstance(plots, dict):
            nested = plots.get("policy") or plots.get("workbook_policy")
            if isinstance(nested, dict):
                out = {
                    str(k).strip(): str(v or "").strip()
                    for k, v in nested.items()
                    if str(k).strip()
                }
            elif nested not in (None, ""):
                out = {"golden_auto": str(nested).strip()}
    ga = str(out.get("golden_auto") or out.get("auto") or "").strip()
    um = str(out.get("ultimate_manual") or out.get("pretty") or "").strip()
    if ga:
        out.setdefault("golden_auto", ga)
        out.setdefault("auto", ga)
    if um:
        out.setdefault("ultimate_manual", um)
        out.setdefault("pretty", um)
    return out


def golden_auto_policy(model: ProductModel | None) -> str:
    mp = workbook_policy_map(model)
    return str(mp.get("golden_auto") or mp.get("auto") or "").strip()


def auto_policy(model: ProductModel | None) -> str:
    return golden_auto_policy(model)


def ultimate_manual_policy(model: ProductModel | None) -> str:
    mp = workbook_policy_map(model)
    return str(mp.get("ultimate_manual") or mp.get("pretty") or "").strip()


def pretty_policy(model: ProductModel | None) -> str:
    mp = workbook_policy_map(model)
    return str(mp.get("pretty") or mp.get("ultimate_manual") or "").strip()


def workbook_policy(model: ProductModel | None) -> str:
    """golden_auto token. Nested map is the schema; string is legacy."""
    return golden_auto_policy(model)


def is_ultimate_path(path: Path | str | None) -> bool:
    """True if this xlsx/folder is the jot / ultimate_manual book."""
    if path is None:
        return False
    text = str(path).replace("\\", "/")
    p = Path(text)
    chunks = [text, p.name, p.stem]
    chunks.extend(str(x) for x in p.parts)
    return any(_ULTIMATE_RX.search(str(c)) for c in chunks if c)


def excel_plots_status(model: ProductModel | None) -> str:
    if model is None:
        return ""
    raw = getattr(model, "excel_plots", None)
    if isinstance(raw, dict):
        return str(raw.get("status") or "").strip()
    return ""


def uses_excel_lock(model: ProductModel | None) -> bool:
    if model is None:
        return False
    plots = normalize_excel_plots(model)
    return golden_auto_policy(model) == WORKBOOK_POLICY or bool(plots)


def policy_errors(model: ProductModel | None) -> list[str]:
    if model is None or not uses_excel_lock(model):
        return []
    errors: list[str] = []
    part = str(getattr(model, "part", "") or "")
    ga = golden_auto_policy(model)
    um = ultimate_manual_policy(model)
    if ga != WORKBOOK_POLICY:
        errors.append(
            f"{part}: workbook_policy.golden_auto must be {WORKBOOK_POLICY!r}, got {ga!r}"
        )
    if um != ULTIMATE_POLICY:
        errors.append(
            f"{part}: workbook_policy.ultimate_manual/pretty must be {ULTIMATE_POLICY!r}, got {um!r}"
        )
    pretty = pretty_policy(model)
    if pretty != PRETTY_POLICY:
        errors.append(
            f"{part}: workbook_policy.pretty must be {PRETTY_POLICY!r} (never auto), got {pretty!r}"
        )
    return errors


def normalize_excel_plots(model: ProductModel | None) -> list[dict[str, str]]:
    """Return [{id, test}, ...] from product_model.excel_plots. Empty if absent."""
    if model is None:
        return []
    raw = getattr(model, "excel_plots", None)
    series_raw: Any = raw
    status = ""
    if isinstance(raw, dict):
        status = str(raw.get("status") or "").strip()
        series_raw = raw.get("series", raw.get("ids", raw.get("plots")))
    out: list[dict[str, str]] = []
    if isinstance(series_raw, dict):
        series_raw = list(series_raw.keys())
    for item in series_raw or []:
        sid = ""
        tes = ""
        if isinstance(item, str):
            sid = item.strip()
        elif isinstance(item, dict):
            sid = str(item.get("id") or item.get("series") or "").strip()
            tes = str(item.get("test") or "").strip()
        if not sid:
            continue
        tes = tes or SERIES_TEST.get(sid, "")
        rec = {"id": sid, "test": tes}
        if status:
            rec["status"] = status
        out.append(rec)
    return out


def bound_series_ids(model: ProductModel | None) -> set[str]:
    return {r["id"] for r in normalize_excel_plots(model) if r.get("id")}


def required_series(model: ProductModel, enabled: list[str] | None) -> list[str]:
    en = {str(x).strip().lower() for x in (enabled or []) if str(x).strip()}
    out: list[str] = []
    if "input_threshold" in en or "vth" in en:
        if model.schmitt:
            out.extend(["vtplus_vs_vcc", "vtminus_vs_vcc", "dvt_vs_vcc"])
        else:
            out.extend(["vih_vs_vcc", "vil_vs_vcc"])
    if "icc" in en:
        out.append("icc_vs_vcc")
    if "delta_icc" in en:
        out.append("delta_icc_vs_vcc")
    if "ii" in en:
        out.append("ii_vs_vcc")
    if "voh" in en:
        out.append("voh_at_ioh")
    if "vol" in en:
        out.append("vol_at_iol")
    if "ioz" in en and model.has_oe():
        out.append("ioz_vs_vcc")
    return out


def series_legal(model: ProductModel, sid: str, enabled: list[str] | None) -> str:
    """Empty if legal. Else a FAIL reason."""
    if sid not in ALLOWED_SERIES:
        return f"excel_plots series {sid!r} is not a card-backed Path B id"
    en = {str(x).strip().lower() for x in (enabled or []) if str(x).strip()}
    if sid == "ioz_vs_vcc" and not model.has_oe():
        return "ioz_vs_vcc only when OE is on the card"
    if sid == "ioz_vs_vcc" and "ioz" not in en:
        return "ioz_vs_vcc only when ioz is enabled"
    if sid == "delta_icc_vs_vcc" and "delta_icc" not in en:
        return "delta_icc_vs_vcc only when delta_icc is enabled+mapped"
    if sid in ("vtplus_vs_vcc", "vtminus_vs_vcc", "dvt_vs_vcc") and not model.schmitt:
        return f"{sid} only when schmitt is true"
    if sid in ("vih_vs_vcc", "vil_vs_vcc") and model.schmitt:
        return f"{sid} is VIH/VIL; Schmitt cards use vtplus/vtminus"
    tes = SERIES_TEST.get(sid, "")
    if tes and tes not in en and not (tes == "input_threshold" and "vth" in en):
        if tes in _PATH_B_DC:
            return f"{sid} test {tes} is not enabled"
    return ""


def detect_series_from_headers(headers: list[str], sheet: str = "") -> set[str]:
    """Regex over known runner headers only. sheet disambiguates ICC/VOH/VOL."""
    found: set[str] = set()
    names = [str(h or "").strip() for h in headers if str(h or "").strip()]
    upper = {n.upper().replace(" ", "") for n in names}
    sh = str(sheet or "").strip()
    for name in names:
        for rx, sid in _HEADER_RX:
            if not rx.match(name):
                continue
            if sid == "icc_vs_vcc" and sh == "DeltaICC":
                found.add("delta_icc_vs_vcc")
                continue
            if sid == "voh_at_ioh" and (sh == "VOL" or "IOL_A" in upper):
                if sh == "VOL" or name.upper() in ("MEASURED", "SPEC_MAX", "VOL"):
                    continue
            if sid == "vol_at_iol" and (sh == "VOH" or "IOH_A" in upper):
                if sh == "VOH" or name.upper() in ("MEASURED", "SPEC_MIN", "VOH"):
                    continue
            found.add(sid)
    if sh == "DeltaICC" or "DELTA_ICC_UA" in upper:
        found.discard("icc_vs_vcc")
        found.add("delta_icc_vs_vcc")
    if sh == "VOH" or "IOH_A" in upper:
        found.discard("vol_at_iol")
    if sh == "VOL" or "IOL_A" in upper:
        found.discard("voh_at_ioh")
    return found


def detect_series_from_rows(rows: list[dict[str, Any]], sheet: str = "") -> set[str]:
    keys: list[str] = []
    for row in rows or []:
        if isinstance(row, dict):
            keys.extend(str(k) for k in row.keys())
    return detect_series_from_headers(keys, sheet)


def binding_errors(
    model: ProductModel,
    enabled: list[str] | None,
    *,
    series_data: set[str] | None = None,
) -> list[str]:
    """FAIL if excel_plots is incomplete or illegal. series_data = detected ids."""
    errors: list[str] = []
    if not uses_excel_lock(model):
        return errors
    bound = bound_series_ids(model)
    need = required_series(model, enabled)
    for sid in need:
        if sid not in bound:
            errors.append(
                f"{model.part}: enabled series {sid} has no excel_plots binding"
            )
    for rec in normalize_excel_plots(model):
        sid = rec.get("id") or ""
        why = series_legal(model, sid, enabled)
        if why:
            errors.append(f"{model.part}: {why}")
    if series_data:
        for sid in sorted(series_data):
            if sid in ALLOWED_SERIES and sid not in bound:
                tes = SERIES_TEST.get(sid, "")
                en = {str(x).strip().lower() for x in (enabled or [])}
                if tes in en or (tes == "input_threshold" and "vth" in en):
                    errors.append(
                        f"{model.part}: series data for {sid} but excel_plots has no binding"
                    )
    errors.extend(policy_errors(model))
    return errors


def default_workbook_name(ctx: Any) -> str:
    model = str(getattr(ctx, "model", "") or getattr(ctx, "part", "") or "Logic").strip()
    package = str(getattr(ctx, "package", "") or "PKG").strip()
    return f"{model}_Lab_Report_{package}.xlsx"


def list_xlsx(workbook_dir: Path) -> list[Path]:
    if not workbook_dir.is_dir():
        return []
    return sorted(p for p in workbook_dir.glob("*.xlsx") if p.is_file() and not p.name.startswith("~$"))


def golden_xlsx(workbook_dir: Path) -> list[Path]:
    """Version golden_auto files only. ultimate_manual jot books are excluded."""
    return [p for p in list_xlsx(workbook_dir) if not is_ultimate_path(p)]


def canonical_workbook_path(ctx: Any, model: ProductModel | None = None) -> Path:
    """golden_auto dest under Version workbook/. Never the jot book."""
    wb_dir = Path(ctx.workbook_dir())
    sm: dict[str, Any] = {}
    load_sm = getattr(ctx, "load_sheet_map", None)
    if callable(load_sm):
        got = load_sm()
        if isinstance(got, dict):
            sm = got
    rel = ""
    block = sm.get("workbook") if isinstance(sm.get("workbook"), dict) else {}
    if isinstance(block, dict):
        rel = str(block.get("path") or "").strip()
    dest: Path | None = None
    if rel:
        candidate = (ctx.manifest_dir() / rel).resolve()
        if not is_ultimate_path(candidate):
            dest = candidate
    existing = golden_xlsx(wb_dir)
    if dest is None:
        if len(existing) == 1:
            dest = existing[0].resolve()
        else:
            dest = (wb_dir / default_workbook_name(ctx)).resolve()
    if is_ultimate_path(dest):
        dest = (wb_dir / default_workbook_name(ctx)).resolve()
    return dest


def orphan_xlsx(ctx: Any, dest: Path) -> list[Path]:
    """Second Version golden book. ultimate_manual files are not golden orphans."""
    dest_r = dest.resolve()
    extras: list[Path] = []
    for p in golden_xlsx(Path(ctx.workbook_dir())):
        if p.resolve() != dest_r:
            extras.append(p)
    return extras


def bind_golden_auto(ctx: Any, model: ProductModel | None = None) -> str:
    """Continue/Open Session / Fill Excel bind to golden_auto Version path only."""
    if model is None:
        from ate.tests.logic.product_model import has_product_model, load_product_model

        key = str(getattr(ctx, "part_key", "") or "").strip().lower()
        if not key or not has_product_model(key):
            return ""
        model = load_product_model(key)
    if not uses_excel_lock(model):
        return ""
    dest = canonical_workbook_path(ctx, model)
    if is_ultimate_path(dest):
        raise UltimateWorkbook(
            f"auto write path is ultimate_manual (never_auto_write): {dest}"
        )
    return str(dest)


def coerce_golden_auto_lab_report(ctx: Any, proposed: str = "") -> str:
    """Overwrite-in-place the Version golden. Refuse ultimate_manual."""
    bound = bind_golden_auto(ctx)
    if not bound:
        return str(proposed or "").strip()
    if proposed and is_ultimate_path(proposed):
        raise UltimateWorkbook(
            f"refusing Path B auto write to ultimate_manual {proposed}"
        )
    return bound


def _dut_n(model: ProductModel, ctx: Any) -> int:
    n = getattr(model, "sample_size", None)
    if n:
        try:
            return max(1, int(n))
        except (TypeError, ValueError):
            pass
    try:
        return max(1, int(getattr(ctx, "sample_size", 1) or 1))
    except (TypeError, ValueError):
        return 1


_SKIP_EXTRA = frozenset({"vector", "fix", "status", "note", "y_expect"})
_ALLOWED_HEADERS = frozenset(
    {
        "DUT",
        "PIN",
        "VCC",
        "VT+",
        "VT-",
        "HYSTERESIS_V",
        "VIH",
        "VIL",
        "VIH_min_V",
        "VIL_max_V",
        "ICC_uA",
        "NEAR_PIN",
        "NEAR_V",
        "OTHERS",
        "VI",
        "II_uA",
        "id",
        "IOH_A",
        "IOL_A",
        "Vref",
        "Measured",
        "Spec_min",
        "Spec_max",
        "Result",
        "mode",
        "pass_mode",
        "VOH",
        "VOL",
        "OE",
        "VOUT",
        "IOZ_uA",
        "DELTA_ICC_uA",
        "vcc_owner",
        "field",
        "value",
    }
)
_IN_PIN_RX = re.compile(r"^IN_[A-Za-z0-9]+$")


def header_allowed(name: str) -> bool:
    k = str(name or "").strip()
    if not k or k in _SKIP_EXTRA:
        return False
    if k in _ALLOWED_HEADERS:
        return True
    return bool(_IN_PIN_RX.match(k))


def invented_headers(headers: list[str]) -> list[str]:
    return [str(h) for h in headers if str(h or "").strip() and not header_allowed(str(h))]


def _headers_for(test_id: str, model: ProductModel, sample_row: dict[str, Any] | None) -> list[str]:
    """Runner row keys only. Do not mint G16-style paste cells."""
    tid = str(test_id or "").strip().lower()
    if tid in ("input_threshold", "vth"):
        if model.schmitt:
            cols = ["DUT", "PIN", "VCC", "VT+", "VT-", "HYSTERESIS_V"]
        else:
            cols = ["DUT", "PIN", "VCC", "VIH", "VIL", "VIH_min_V", "VIL_max_V"]
    elif tid == "icc":
        cols = ["DUT", "VCC", "ICC_uA"] + [f"IN_{p}" for p in icc_pins(model)]
    elif tid == "delta_icc":
        cols = ["DUT", "VCC", "NEAR_PIN", "NEAR_V", "OTHERS", "ICC_uA"]
    elif tid == "ii":
        cols = ["DUT", "PIN", "VCC", "VI", "II_uA"]
    elif tid == "voh":
        cols = ["DUT", "id", "VCC", "IOH_A", "Vref", "Measured", "Spec_min", "Result", "mode", "pass_mode"]
    elif tid == "vol":
        cols = ["DUT", "id", "VCC", "IOL_A", "Vref", "Measured", "Spec_max", "Result", "mode", "pass_mode"]
    elif tid == "ioz":
        cols = ["DUT", "VCC", "OE", "VOUT", "IOZ_uA"]
    else:
        cols = ["DUT", "VCC"]
    extra: list[str] = []
    if isinstance(sample_row, dict):
        for key in sample_row.keys():
            k = str(key)
            if k not in cols and header_allowed(k):
                extra.append(k)
    return cols + extra


def _template_body(test_id: str, model: ProductModel, ctx: Any) -> list[dict[str, Any]]:
    """Empty DUT x VCC (x PIN) rows. No invented measurements."""
    tid = str(test_id or "").strip().lower()
    dut_n = _dut_n(model, ctx)
    vccs: list[Any] = list(model.vcc_list) or [None]
    pins: list[Any] = list(model.logic_inputs) or [None]
    body: list[dict[str, Any]] = []
    per_pin = tid in ("input_threshold", "vth", "ii")
    for dut in range(1, dut_n + 1):
        if per_pin:
            for pin in pins:
                for vcc in vccs:
                    rec: dict[str, Any] = {"DUT": dut}
                    if pin:
                        rec["PIN"] = pin
                    if vcc is not None:
                        rec["VCC"] = vcc
                    body.append(rec)
        else:
            for vcc in vccs:
                rec = {"DUT": dut}
                if vcc is not None:
                    rec["VCC"] = vcc
                body.append(rec)
    return body


def _sheet_tests(enabled: list[str]) -> dict[str, list[str]]:
    """lab_sheet -> [test ids] for Path B DC only."""
    grouped: dict[str, list[str]] = {}
    for tid in enabled or []:
        t = str(tid).strip().lower()
        if t not in _PATH_B_DC:
            continue
        sheet = _SHEET.get(t)
        if not sheet:
            continue
        grouped.setdefault(sheet, [])
        if t not in grouped[sheet]:
            grouped[sheet].append(t)
    return grouped


def _steps_rows(report: dict[str, Any] | None, test_ids: list[str]) -> list[dict[str, Any]]:
    want = {str(t).strip().lower() for t in test_ids}
    rows: list[dict[str, Any]] = []
    if not isinstance(report, dict):
        return rows
    for step in report.get("steps") or []:
        if not isinstance(step, dict):
            continue
        tid = str(step.get("test_id") or "").strip().lower()
        if tid not in want:
            continue
        try:
            dut = max(1, int(step.get("dut") or 1))
        except (TypeError, ValueError):
            dut = 1
        data = step.get("data") if isinstance(step.get("data"), dict) else {}
        raw_rows = data.get("rows") if isinstance(data, dict) else None
        if isinstance(raw_rows, list) and raw_rows:
            for rec in raw_rows:
                if not isinstance(rec, dict):
                    continue
                item = dict(rec)
                item.setdefault("DUT", dut)
                rows.append(item)
            continue
        meas = step.get("measurements") if isinstance(step.get("measurements"), list) else []
        rec: dict[str, Any] = {"DUT": dut}
        for m in meas:
            if not isinstance(m, dict):
                continue
            mid = str(m.get("id") or "").strip()
            if mid:
                rec[mid] = m.get("value")
        if len(rec) > 1:
            rows.append(rec)
    return rows


def _cell_dump(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, (int, float, str, bool)):
        return val
    try:
        return json.dumps(val, ensure_ascii=True, sort_keys=True)
    except TypeError:
        return str(val)


def _write_setup(ws: Any, model: ProductModel) -> None:
    ws.title = "Setup"
    rows = [
        ("part", model.part),
        ("status", model.status),
        ("truth_table_status", model.truth_table_status),
        ("isolation_status", model.isolation_status),
        ("schmitt", model.schmitt),
        ("oe", model.oe_mode),
        ("workbook_policy", workbook_policy_map(model) or {"golden_auto": WORKBOOK_POLICY}),
        ("golden_auto", golden_auto_policy(model) or WORKBOOK_POLICY),
        ("auto", golden_auto_policy(model) or WORKBOOK_POLICY),
        ("ultimate_manual", ultimate_manual_policy(model) or ULTIMATE_POLICY),
        ("pretty", pretty_policy(model) or PRETTY_POLICY),
        ("vcc_list", list(model.vcc_list)),
        ("vcc_grid", dict(model.vcc_grid or {})),
        ("pass_mode", dict(model.pass_mode)),
        ("excel_plots", [r.get("id") for r in normalize_excel_plots(model)]),
        ("wire_map", dict(model.wire_map or {})),
        ("sample_size", model.sample_size),
    ]
    ws["A1"] = "field"
    ws["B1"] = "value"
    for i, (key, val) in enumerate(rows, start=2):
        ws.cell(i, 1, key)
        ws.cell(i, 2, _cell_dump(val))
    leftover = int(getattr(ws, "max_row", 1) or 1)
    if leftover > len(rows) + 1:
        ws.delete_rows(len(rows) + 2, leftover - (len(rows) + 1))


def _replace_sheet(wb: Any, name: str, index: int | None = None) -> Any:
    """Overwrite a tab in the same workbook. Never creates a second xlsx."""
    names = list(wb.sheetnames)
    if name in names:
        if len(names) == 1:
            ws = wb[name]
            ws.title = name
            charts = getattr(ws, "_charts", None)
            if charts is not None:
                charts.clear()
            max_r = int(getattr(ws, "max_row", 1) or 1)
            max_c = int(getattr(ws, "max_column", 1) or 1)
            if max_r:
                ws.delete_rows(1, max_r)
            if max_c:
                for col in range(1, max_c + 1):
                    ws.cell(1, col, None)
            return ws
        idx = names.index(name)
        del wb[name]
        return wb.create_sheet(name, idx)
    if index is None:
        return wb.create_sheet(name)
    return wb.create_sheet(name, index)


def _col_index(headers: list[str], name: str) -> int | None:
    want = str(name).strip().lower()
    for i, h in enumerate(headers):
        if str(h).strip().lower() == want:
            return i + 1
    return None


def _has_numeric(ws: Any, col: int, n_rows: int) -> bool:
    n = 0
    for r in range(2, n_rows + 2):
        val = ws.cell(r, col).value
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            n += 1
    return n >= 2


def _add_charts(ws: Any, headers: list[str], n_rows: int, series_ids: list[str]) -> int:
    if n_rows < 2:
        return 0
    from openpyxl.chart import LineChart, Reference, ScatterChart
    try:
        from openpyxl.chart import Series
    except ImportError:
        from openpyxl.chart.series import Series

    charts = getattr(ws, "_charts", None)
    if charts is not None:
        charts.clear()
    added = 0
    anchor_row = 2
    for sid in series_ids:
        x_name = _X_FOR_SERIES.get(sid, "VCC")
        x_col = _col_index(headers, x_name)
        y_names = _Y_FOR_SERIES.get(sid) or ()
        y_cols = [c for c in (_col_index(headers, y) for y in y_names) if c]
        if not x_col or not y_cols:
            continue
        if not _has_numeric(ws, x_col, n_rows):
            continue
        y_ok = [c for c in y_cols if _has_numeric(ws, c, n_rows)]
        if not y_ok:
            continue
        scatter = sid in ("voh_at_ioh", "vol_at_iol")
        chart = ScatterChart() if scatter else LineChart()
        chart.title = sid
        chart.y_axis.title = str(y_names[0])
        chart.x_axis.title = x_name
        if scatter:
            xvalues = Reference(ws, min_col=x_col, min_row=2, max_row=n_rows + 1)
            for yc in y_ok:
                yvalues = Reference(ws, min_col=yc, min_row=1, max_row=n_rows + 1)
                chart.series.append(Series(yvalues, xvalues, title_from_data=True))
        else:
            cats = Reference(ws, min_col=x_col, min_row=2, max_row=n_rows + 1)
            chart.set_categories(cats)
            for yc in y_ok:
                data = Reference(ws, min_col=yc, min_row=1, max_row=n_rows + 1)
                chart.add_data(data, titles_from_data=True)
        if not chart.series:
            continue
        ws.add_chart(chart, f"H{anchor_row}")
        anchor_row += 16
        added += 1
    return added


def _write_test_sheet(
    ws: Any,
    *,
    sheet: str,
    test_ids: list[str],
    model: ProductModel,
    ctx: Any,
    report: dict[str, Any] | None,
    bound: set[str],
) -> int:
    ws.title = sheet
    primary = test_ids[0]
    rows = _steps_rows(report, test_ids)
    sample = rows[0] if rows else None
    headers = _headers_for(primary, model, sample)
    bad = invented_headers(headers)
    if bad:
        raise RuntimeError(f"invented columns {bad} (card-backed runner ids only)")
    for c, name in enumerate(headers, start=1):
        ws.cell(1, c, name)
    body = rows or _template_body(primary, model, ctx)
    for r, rec in enumerate(body, start=2):
        for c, name in enumerate(headers, start=1):
            ws.cell(r, c, _cell_dump(rec.get(name)))
    detected = detect_series_from_headers(headers, sheet)
    plot_ids = [sid for sid in required_series(model, test_ids) if sid in bound]
    plot_ids.extend(sid for sid in sorted(detected & bound) if sid not in plot_ids)
    if rows:
        return _add_charts(ws, headers, len(body), plot_ids)
    return 0


def write_path_b_workbook(
    *,
    ctx: Any,
    model: ProductModel,
    report: dict[str, Any] | None = None,
    enabled: list[str] | None = None,
) -> dict[str, Any]:
    """Create or overwrite the one Version golden_auto xlsx. Never writes ultimate_manual."""
    errors = policy_errors(model)
    if errors:
        raise RuntimeError("; ".join(errors))
    if enabled is None:
        from ate.fixture.modes import enabled_tests_for_part

        key = str(getattr(ctx, "part_key", "") or "").strip().lower()
        enabled = list(enabled_tests_for_part(key) or [])
    dest = canonical_workbook_path(ctx, model)
    if is_ultimate_path(dest):
        raise UltimateWorkbook(
            f"auto write path == ultimate_manual (never_auto_write): {dest}"
        )
    if dest.name.endswith("_filled.xlsx") or dest.stem.endswith("_filled"):
        raise OrphanWorkbook(f"refusing orphan path {dest}")
    wb_dir = Path(ctx.workbook_dir())
    wb_dir.mkdir(parents=True, exist_ok=True)
    extras = orphan_xlsx(ctx, dest)
    if extras:
        raise OrphanWorkbook(
            f"workbook/ has extra golden xlsx {sorted(p.name for p in extras)}; "
            f"golden_auto dest={dest.name}"
        )
    from openpyxl import Workbook, load_workbook

    wb = None
    created = not dest.is_file()
    try:
        if dest.is_file():
            wb = load_workbook(dest)
        else:
            wb = Workbook()
            active = wb.active
            if active is not None and active.title in ("Sheet", "Sheet1"):
                active.title = "Setup"
        grouped = _sheet_tests(enabled or [])
        if not model.has_oe():
            grouped.pop("IOZ", None)
        if "delta_icc" not in {str(t).lower() for t in (enabled or [])}:
            grouped.pop("DeltaICC", None)
        known = set(_SHEET.values())
        for name in list(wb.sheetnames):
            if name in known and name not in grouped and len(wb.sheetnames) > 1:
                del wb[name]
        setup = _replace_sheet(wb, "Setup", 0)
        _write_setup(setup, model)
        bound = bound_series_ids(model)
        plots = 0
        for sheet, tids in grouped.items():
            if sheet == "IOZ" and not model.has_oe():
                continue
            ws = _replace_sheet(wb, sheet)
            plots += _write_test_sheet(
                ws,
                sheet=sheet,
                test_ids=tids,
                model=model,
                ctx=ctx,
                report=report,
                bound=bound,
            )
        try:
            wb.save(dest)
        except PermissionError as exc:
            raise OrphanWorkbook(
                f"xlsx locked at {dest}; golden_auto refuses a second book"
            ) from exc
        extras_after = orphan_xlsx(ctx, dest)
        if extras_after:
            raise OrphanWorkbook(
                f"write created orphan golden xlsx {sorted(p.name for p in extras_after)}"
            )
        return {
            "filled": 1,
            "status": "ok",
            "excel": str(dest),
            "created": created,
            "target": "golden_auto",
            "policy": workbook_policy_map(model),
            "plots": plots,
        }
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
